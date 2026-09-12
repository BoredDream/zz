"""确认附件4 与 result4-2/4-3 模板的结构。"""
from __future__ import annotations

from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "problem" / "data"

p4 = pd.read_excel(DATA / "附件4.xlsx", header=0)
print("附件4 表头:", list(p4.columns[:4]), "...", list(p4.columns[-2:]))
arr = p4.iloc[:, 1:].to_numpy(float)
print("附件4 形状:", arr.shape, " 价格 min/max/mean = %.4f / %.4f / %.4f" % (arr.min(), arr.max(), arr.mean()))
print("附件4 首行前 6:", np.round(arr[0, :6], 4), " 末 3:", np.round(arr[0, -3:], 4))
print("附件4 日内均价 min/max:", np.round(arr.mean(1).min(), 4), np.round(arr.mean(1).max(), 4))
print("附件4 逐时段均值 min/max:", np.round(arr.mean(0).min(), 4), np.round(arr.mean(0).max(), 4))

for name in ("result4-2.xlsx", "result4-3.xlsx"):
    wb = openpyxl.load_workbook(DATA / "附件5" / name)
    print(f"\n[{name}] sheets={wb.sheetnames}")
    for s in wb.sheetnames:
        ws = wb[s]
        hdr = [ws.cell(1, c).value for c in range(1, min(ws.max_column, 5) + 1)]
        tail = [ws.cell(1, c).value for c in range(max(1, ws.max_column - 1), ws.max_column + 1)]
        print(f"   {s:<8} dims={ws.dimensions:<12} 表头={hdr} ... {tail}")
        if s == "充放电量":
            for r in range(2, min(ws.max_row, 8) + 1):
                print("      r%d" % r, [ws.cell(r, c).value for c in range(1, 7)])
