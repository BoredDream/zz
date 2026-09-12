from __future__ import annotations

import argparse, csv, itertools, json
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any
import numpy as np
from openpyxl import load_workbook
from scipy.optimize import linprog
from scipy.sparse import coo_matrix
from q2_solver import Config, choose_weight, emergency_segments, energy_score, file_sha256, point_forecast

ISSUE_INTERVALS=(0,36,72,108); ISSUE_LABELS=("0:00","6:00","12:00","18:00")
TARGET_DATES={date(2025,3,20),date(2025,6,21),date(2025,9,23),date(2025,12,21)}; DELIVERY_START=date(2025,2,1)

@dataclass(frozen=True)
class Q3Config(Config):
    pv_interpolation: str="linear_from_latest_observation_to_hourly_endpoints"

def parse_date(v: Any)->date: return v.date() if isinstance(v,datetime) else datetime.strptime(str(v),"%Y-%m-%d").date()

def load_q3_inputs(data_dir:Path,cfg:Q3Config)->dict[str,Any]:
    pp,ap,fp=data_dir/"附件1.xlsx",data_dir/"附件2.xlsx",data_dir/"附件3.xlsx"
    pr=list(load_workbook(pp,read_only=True,data_only=True).active.iter_rows(min_row=2,max_row=145,min_col=1,max_col=4,values_only=True))
    raw_price=np.array([float(r[1]) for r in pr]); cold_l_raw=np.array([float(r[2]) for r in pr]); cold_p_raw=np.array([float(r[3]) for r in pr])
    wb=load_workbook(ap,read_only=True,data_only=True)
    lr=list(wb["小区负载"].iter_rows(min_row=2,max_row=366,min_col=1,max_col=145,values_only=True)); pv=list(wb["光伏发电实际功率"].iter_rows(min_row=2,max_row=366,min_col=1,max_col=145,values_only=True))
    dates=[r[0].date() for r in lr]; expected=[date(2025,1,1)+timedelta(days=i) for i in range(365)]
    if dates!=expected or [r[0].date() for r in pv]!=expected: raise ValueError("附件2日期不连续或两表日期不一致")
    lraw=np.array([[float(v) for v in r[1:]] for r in lr]); praw=np.array([[float(v) for v in r[1:]] for r in pv])
    lnat=np.full_like(lraw,np.nan); pnat=np.full_like(praw,np.nan); lnat[:,1:]=lraw[:,:143]; pnat[:,1:]=praw[:,:143]; lnat[1:,0]=lraw[:-1,143]; pnat[1:,0]=praw[:-1,143]
    cold_l=np.r_[cold_l_raw[-1],cold_l_raw[:-1]]; cold_p=np.r_[cold_p_raw[-1],cold_p_raw[:-1]]
    rows=list(load_workbook(fp,read_only=True,data_only=True).active.iter_rows(min_row=2,max_row=1461,min_col=1,max_col=26,values_only=True))
    if len(rows)!=1460: raise ValueError("附件3应有365天×4个发布时刻")
    forecasts=np.zeros((365,4,24)); fdates=[]
    for d in range(365):
        group=rows[4*d:4*d+4]; fdates.append(parse_date(group[0][0]))
        if [str(r[1]) for r in group]!=list(ISSUE_LABELS): raise ValueError(f"附件3第{d+1}天发布时刻不完整")
        forecasts[d]=np.array([[float(v) for v in r[2:26]] for r in group])
    if fdates!=expected: raise ValueError("附件3日期不连续")
    feb=dates.index(DELIVERY_START)
    return {"dates":dates,"price_natural":np.r_[raw_price[-1],raw_price[:-1]],"price_template":raw_price,"load_actual_kw":lnat,"pv_actual_kw":pnat,"net_actual_kwh":(lnat-pnat)*cfg.dt_hours,"cold_load_kw":cold_l,"cold_pv_kw":cold_p,"pv_forecast_kw":forecasts,
      "mapping_audit":{"actual_midnights_checked":364,"actual_midnight_max_error_kw":float(max(np.nanmax(abs(lnat[1:,0]-lraw[:-1,143])),np.nanmax(abs(pnat[1:,0]-praw[:-1,143])))),"feb1_midnight_net_kwh":float((lnat[feb,0]-pnat[feb,0])*cfg.dt_hours),"template_row_coverage":"00:10至次日00:00；自然日统计为00:00至24:00"},"input_hashes":{p.name:file_sha256(p) for p in (pp,ap,fp)}}

def absolute_profile(values,cold,day,weight): return np.r_[point_forecast(values,cold,day,day,weight),point_forecast(values,cold,day+1,day,weight)[0]]
def pv_curve(inputs,day,source_i):
    issue=ISSUE_INTERVALS[source_i]; anchor=float(inputs["pv_actual_kw"][day-1,143]) if issue==0 and day>0 else (float(inputs["cold_pv_kw"][0]) if issue==0 else float(inputs["pv_actual_kw"][day,issue-1]))
    return np.interp(np.arange(145-issue),np.arange(0,145,6),np.r_[anchor,inputs["pv_forecast_kw"][day,source_i]])
def forecast_net(inputs,day,target_i,source_i,weight,cfg):
    target,source=ISSUE_INTERVALS[target_i],ISSUE_INTERVALS[source_i]
    if source>target: raise ValueError("预报源时刻晚于决策时刻")
    return (absolute_profile(inputs["load_actual_kw"],inputs["cold_load_kw"],day,weight)[target:]-pv_curve(inputs,day,source_i)[target-source:])*cfg.dt_hours
def actual_horizon(inputs,day,issue_i):
    issue=ISSUE_INTERVALS[issue_i]; return np.r_[inputs["net_actual_kwh"][day,issue:],inputs["net_actual_kwh"][day+1,0]]
def choose_residual_window(residuals,latest,cfg):
    best,bscore=cfg.residual_window_candidates[0],np.inf
    for window in cfg.residual_window_candidates:
        scores=[]
        for target in range(max(1,latest-cfg.distribution_validation_days+1),latest+1):
            train=[j for j in range(max(1,target-window),target) if j in residuals]
            if train: scores.append(energy_score(np.stack([residuals[j] for j in train]),residuals[target]))
        score=float(np.mean(scores)) if scores else np.inf
        if score<bscore: best,bscore=window,score
    return best
def make_scenarios(inputs,day,target_i,source_i,weight,cfg):
    center=forecast_net(inputs,day,target_i,source_i,weight,cfg); latest=day-2 if target_i==0 else day-1
    residuals={j:actual_horizon(inputs,j,target_i)-forecast_net(inputs,j,target_i,source_i,weight,cfg) for j in range(1,max(1,latest+1))}
    window=choose_residual_window(residuals,latest,cfg); selected=[j for j in range(max(1,latest-window+1),latest+1) if j in residuals]
    return (center[None,:] if not selected else center[None,:]+np.stack([residuals[j] for j in selected])),center,window

def solve_common(scenarios,prices,initial_soc,terminal,cfg,fixed_first=None,reference=None,mode="plan"):
    """All scenarios share C,D,S; only emergency and surplus are recourse."""
    W,H=scenarios.shape; G=H-1 if mode=="plan" else H; inc0=G; dec0=inc0+(0 if mode=="plan" else G); common=G if mode=="plan" else 3*G
    c0,d0,s0=common,common+H,common+2*H; rec=s0+H+1; n=rec+2*W*H; obj=np.zeros(n)
    if mode=="plan": obj[:G]=prices[1:]
    elif mode=="per_submission": obj[inc0:inc0+G]=1.5*prices; obj[dec0:dec0+G]=.5*prices
    elif mode=="final_refund": obj[:G]=prices; obj[inc0:inc0+G]=.5*prices; obj[dec0:dec0+G]=.5*prices
    else: raise ValueError(mode)
    obj[c0:c0+H]=cfg.throughput_penalty; obj[d0:d0+H]=cfg.throughput_penalty; obj[s0+H]=-terminal
    for w in range(W): obj[rec+2*w*H:rec+(2*w+1)*H]=cfg.emergency_price_multiple*prices/W
    rr=[];cc=[];vv=[];rhs=[];row=0
    if mode!="plan":
        if reference is None or len(reference)!=G: raise ValueError("reference长度错误")
        for h in range(G):
            for col,val in ((h,1),(inc0+h,-1),(dec0+h,1)): rr.append(row);cc.append(col);vv.append(val)
            rhs.append(float(reference[h]));row+=1
    for w in range(W):
        e0=rec+2*w*H;u0=e0+H
        for h in range(H):
            const=0.0
            if mode=="plan":
                if h==0: const=float(fixed_first)
                else: rr.append(row);cc.append(h-1);vv.append(1)
            else: rr.append(row);cc.append(h);vv.append(1)
            for col,val in ((c0+h,-1),(d0+h,1),(e0+h,1),(u0+h,-1)): rr.append(row);cc.append(col);vv.append(val)
            rhs.append(float(scenarios[w,h]-const));row+=1
    for h in range(H):
        for col,val in ((s0+h+1,1),(s0+h,-1),(c0+h,-cfg.eta_charge),(d0+h,1/cfg.eta_discharge)): rr.append(row);cc.append(col);vv.append(val)
        rhs.append(0.0);row+=1
    rr.append(row);cc.append(s0);vv.append(1);rhs.append(initial_soc);row+=1
    mat=coo_matrix((vv,(rr,cc)),shape=(row,n)).tocsr(); bounds=[(0,None)]*G
    if mode!="plan": bounds += [(0,None)]*(2*G)
    bounds += [(0,cfg.interval_limit_kwh)]*(2*H)+[(cfg.soc_min_kwh,cfg.soc_max_kwh)]*(H+1)+[(0,None)]*(2*W*H)
    res=linprog(obj,A_eq=mat,b_eq=np.array(rhs),bounds=bounds,method="highs")
    if not res.success: raise RuntimeError(res.message)
    out={"grid":res.x[:G],"charge":res.x[c0:c0+H],"discharge":res.x[d0:d0+H],"soc":res.x[s0:s0+H+1]}
    if mode!="plan": out.update(increase=res.x[inc0:inc0+G],decrease=res.x[dec0:dec0+G])
    return out

def execute_fixed(actual,grid,charge,discharge,soc0,cfg):
    soc=np.zeros(len(actual)+1);soc[0]=soc0
    for t in range(len(actual)): soc[t+1]=soc[t]+cfg.eta_charge*charge[t]-discharge[t]/cfg.eta_discharge
    gap=actual-(grid+discharge-charge)
    return {"charge":charge.copy(),"discharge":discharge.copy(),"emergency":np.maximum(gap,0),"surplus":np.maximum(-gap,0),"soc":soc}
def policy_label(updates,forecast_mode="latest",settlement="per_submission"):
    x="none" if not updates else "+".join(str(ISSUE_INTERVALS[i]//6) for i in updates)
    if forecast_mode=="frozen0": x+="_frozen0"
    if settlement=="final_refund": x+="_refund"
    return x
def simulate_policy(inputs,cfg,updates,forecast_mode="latest",settlement="per_submission"):
    label=policy_label(updates,forecast_mode,settlement);net=inputs["net_actual_kwh"].copy();net[0,0]=(inputs["cold_load_kw"][0]-inputs["cold_pv_kw"][0])*cfg.dt_hours
    price=inputs["price_natural"];terminal=cfg.eta_discharge*float(price.min());soc_now=cfg.initial_soc_kwh;prev_plan=max(float(net[0,0]),0);prev_active=prev_plan;prev_penalty=0.; records=[];releases=[]
    for day,dt in enumerate(inputs["dates"]):
        weight=choose_weight(inputs["load_actual_kw"],inputs["cold_load_kw"],day,cfg);sc0,center0,w0=make_scenarios(inputs,day,0,0,weight,cfg)
        sol0=solve_common(sc0,np.r_[price,price[0]],soc_now,terminal,cfg,fixed_first=prev_active);plan=sol0["grid"].copy();active=plan.copy();penalty=np.zeros(144);C=sol0["charge"].copy();D=sol0["discharge"].copy()
        trace={k:[] for k in ("charge","discharge","emergency","surplus")};trace["soc"]=[soc_now];grids=[];metrics=[(0,center0[:36],net[day,:36])];windows={"0:00":w0};counts={"0:00":len(sc0)}
        releases.append({"policy":label,"date":dt.isoformat(),"issue":"0:00","forecast_source":"0:00","information_cutoff":"previous_day_23:50","reference_kind":"new_plan","window_days":w0,"scenario_count":len(sc0),"increase_kwh":0.,"decrease_kwh":0.,"adjustment_fee_yuan":0.,"contract_kwh":list(map(float,plan)),"common_charge_kwh":list(map(float,C)),"common_discharge_kwh":list(map(float,D))})
        for issue_i,issue in enumerate(ISSUE_INTERVALS):
            if issue_i>0 and issue_i in updates:
                source=issue_i if forecast_mode=="latest" else 0;sc,center,win=make_scenarios(inputs,day,issue_i,source,weight,cfg);p=np.r_[price[issue:],price[0]];ref=(active if settlement=="per_submission" else plan)[issue-1:].copy();old=active[issue-1:].copy()
                sol=solve_common(sc,p,soc_now,terminal,cfg,reference=ref,mode=settlement);active[issue-1:]=sol["grid"];C[issue:]=sol["charge"];D[issue:]=sol["discharge"]
                inc=np.maximum(sol["grid"]-old,0);dec=np.maximum(old-sol["grid"],0);fee=1.5*p*inc+.5*p*dec if settlement=="per_submission" else np.zeros_like(inc)
                if settlement=="per_submission": penalty[issue-1:]+=fee
                releases.append({"policy":label,"date":dt.isoformat(),"issue":ISSUE_LABELS[issue_i],"forecast_source":ISSUE_LABELS[source],"information_cutoff":f"{issue//6-1:02d}:50","reference_kind":"previous_active" if settlement=="per_submission" else "original_plan","window_days":win,"scenario_count":len(sc),"increase_kwh":float(inc.sum()),"decrease_kwh":float(dec.sum()),"adjustment_fee_yuan":float(fee.sum()),"contract_kwh":list(map(float,sol["grid"])),"common_charge_kwh":list(map(float,sol["charge"])),"common_discharge_kwh":list(map(float,sol["discharge"]))})
                metrics.append((issue_i,center[:min(36,144-issue)],net[day,issue:issue+36]));windows[ISSUE_LABELS[issue_i]]=win;counts[ISSUE_LABELS[issue_i]]=len(sc)
            block=min(36,144-issue);g=np.r_[prev_active,active[:35]] if issue==0 else active[issue-1:issue-1+block]
            run=execute_fixed(net[day,issue:issue+block],g,C[issue:issue+block],D[issue:issue+block],soc_now,cfg);grids.extend(map(float,g))
            for k in ("charge","discharge","emergency","surplus"): trace[k].extend(map(float,run[k]))
            trace["soc"].extend(map(float,run["soc"][1:]));soc_now=float(run["soc"][-1])
        base=np.r_[prev_plan,plan[:143]];act=np.r_[prev_active,active[:143]];pen=np.r_[prev_penalty,penalty[:143]];em=np.array(trace["emergency"]);ecost=float(cfg.emergency_price_multiple*price@em)
        scost=float(price@base+pen.sum()) if settlement=="per_submission" else float(price@(act+.5*np.abs(act-base)))
        records.append({"date":dt,"policy":label,"weight":weight,"windows":windows,"scenario_counts":counts,"plan":plan,"adjusted":active,"base_natural":base,"adjusted_natural":act,"penalty_natural":pen,"penalty_template":penalty.copy(),"grid":np.array(grids),"charge":np.array(trace["charge"]),"discharge":np.array(trace["discharge"]),"emergency":em,"surplus":np.array(trace["surplus"]),"soc":np.array(trace["soc"]),"settlement_cost":scost,"emergency_cost":ecost,"issue_metrics":metrics})
        prev_plan=float(plan[-1]);prev_active=float(active[-1]);prev_penalty=float(penalty[-1])
        if (day+1)%100==0: print(f"{label}: {day+1}/365",flush=True)
    return records,releases

def block_sums(v): return [float(v[24*i:24*(i+1)].sum()) for i in range(6)]
def clock(i): return "24:00" if i==144 else f"{i//6:02d}:{(i%6)*10:02d}"
def write_csv(path,rows):
    with path.open("w",newline="",encoding="utf-8-sig") as f: w=csv.DictWriter(f,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)
def totals(records,start):
    rs=records[start:];s=sum(r["settlement_cost"] for r in rs);e=sum(r["emergency_cost"] for r in rs)
    return {"settlement_cost_yuan":float(s),"emergency_cost_yuan":float(e),"total_cost_yuan":float(s+e),"emergency_kwh":float(sum(r["emergency"].sum() for r in rs))}
def detail_rows(records,inputs,start):
    out=[];p=inputs["price_natural"]
    for d,r in enumerate(records[start:],start=start):
        for t in range(144): out.append({"policy":r["policy"],"date":r["date"].isoformat(),"time_start":clock(t),"price_yuan_per_kwh":float(p[t]),"initial_plan_kwh":float(r["base_natural"][t]),"final_contract_kwh":float(r["adjusted_natural"][t]),"adjustment_fee_yuan":float(r["penalty_natural"][t]),"charge_kwh":float(r["charge"][t]),"discharge_kwh":float(r["discharge"][t]),"soc_start_kwh":float(r["soc"][t]),"soc_end_kwh":float(r["soc"][t+1]),"actual_net_kwh":float(inputs["net_actual_kwh"][d,t]),"emergency_kwh":float(r["emergency"][t]),"surplus_kwh":float(r["surplus"][t])})
    return out

def assemble(primary,policies,releases,inputs,cfg):
    start=inputs["dates"].index(DELIVERY_START);delivered=primary[start:];price=inputs["price_natural"];tp=inputs["price_template"];ct={k:totals(v,start) for k,v in policies.items()};daily=[];intervals=detail_rows(primary,inputs,start);days=[];errors={x:[] for x in ISSUE_LABELS};mb=ms=0.
    for idx,r in enumerate(delivered,start=start):
        actual=inputs["net_actual_kwh"][idx];bal=r["grid"]+r["emergency"]+r["discharge"]-r["charge"]-r["surplus"]-actual;sr=r["soc"][1:]-r["soc"][:-1]-cfg.eta_charge*r["charge"]+r["discharge"]/cfg.eta_discharge;mb=max(mb,float(abs(bal).max()));ms=max(ms,float(abs(sr).max()))
        for ii,f,o in r["issue_metrics"]: errors[ISSUE_LABELS[ii]].extend(map(float,o-f))
        cs,ds=block_sums(r["charge"]),block_sums(r["discharge"]);blocks=[{"time_range":f"{4*i}:00-{4*(i+1)}:00","charge_kwh":cs[i],"discharge_kwh":ds[i]} for i in range(6)];template_pen=r["penalty_template"]
        day={"date":r["date"].isoformat(),"plan_kwh":list(map(float,r["plan"])),"adjusted_kwh":list(map(float,r["adjusted"])),"plan_total_kwh":float(r["plan"].sum()),"plan_cost_yuan":float(tp@r["plan"]),"adjusted_total_kwh":float(r["adjusted"].sum()),"adjusted_settlement_yuan":float(tp@r["plan"]+template_pen.sum()),"storage_blocks":blocks,"soc_start_kwh":float(r["soc"][0]),"soc_end_kwh":float(r["soc"][-1]),"emergency_segments":emergency_segments(r["emergency"]),"emergency_total_kwh":float(r["emergency"].sum()),"emergency_cost_yuan":float(r["emergency_cost"])};days.append(day)
        daily.append({"policy":r["policy"],"date":day["date"],"plan_cost_yuan":day["plan_cost_yuan"],"adjustment_fee_yuan":float(r["penalty_natural"].sum()),"settlement_cost_yuan":r["settlement_cost"],"emergency_kwh":day["emergency_total_kwh"],"emergency_cost_yuan":r["emergency_cost"],"total_cost_yuan":r["settlement_cost"]+r["emergency_cost"],"soc_start_kwh":day["soc_start_kwh"],"soc_end_kwh":day["soc_end_kwh"]})
    fixed_ref=float(sum(price@(r["adjusted_natural"]+.5*np.abs(r["adjusted_natural"]-r["base_natural"]))+r["emergency_cost"] for r in delivered));acc={k:{"mae_kwh_per_interval":float(np.mean(np.abs(v))),"rmse_kwh_per_interval":float(np.sqrt(np.mean(np.array(v)**2))),"observations":len(v)} for k,v in errors.items()}
    C=np.concatenate([r["charge"] for r in delivered]);D=np.concatenate([r["discharge"] for r in delivered]);E=np.concatenate([r["emergency"] for r in delivered]);S=np.concatenate([r["soc"] for r in delivered])
    summary={"model":"causal_rolling_SAA_MPC_with_common_storage_actions","strict_multistage_optimal":False,"period":{"start":"2025-02-01","end":"2025-12-31","days":334},"time_convention":"区间起点；模板行00:10至次日00:00；自然日00:00至24:00","primary_settlement_assumption":"原计划费用保留；每次提交相对当前合同逐笔计费：上调1.5p、下调0.5p；需人工确认","alternative_settlement_assumption":"最终净额退款口径：pA+0.5p|A-G|；需人工确认","totals_natural_day":ct["6+12+18"],"policy_comparison":ct,"fixed_quantity_refund_resettlement_total_cost_yuan":fixed_ref,"forecast_accuracy_next_block":acc,"mapping_audit":inputs["mapping_audit"],"checks":{"max_energy_balance_residual_kwh":mb,"max_soc_residual_kwh":ms,"soc_min_observed_kwh":float(S.min()),"soc_max_observed_kwh":float(S.max()),"max_interval_storage_energy_kwh":float(max(C.max(),D.max())),"simultaneous_charge_discharge_intervals":int(np.count_nonzero((C>1e-7)&(D>1e-7))),"emergency_while_charging_intervals":int(np.count_nonzero((E>1e-7)&(C>1e-7))),"cross_day_soc_continuity":all(abs(delivered[i]["soc"][-1]-delivered[i+1]["soc"][0])<1e-7 for i in range(len(delivered)-1))},"target_dates":{d["date"]:d for d in days if date.fromisoformat(d["date"]) in TARGET_DATES},"input_hashes_sha256":inputs["input_hashes"]}
    return summary,{"metadata":{"time_convention":summary["time_convention"],"settlement":summary["primary_settlement_assumption"],"input_hashes_sha256":inputs["input_hashes"]},"days":days},{"daily":daily,"interval":intervals,"releases":releases}

def fmt(x): return f"{x:,.4f}"
def build_report(s):
    t=s["totals_natural_day"];c=s["policy_comparison"];base=c["none"]["total_cost_yuan"]
    L=["# 第三问修正模型计算报告","","## 1. 结论","","采用因果滚动SAA/MPC。每个发布时刻的全部情景共享同一条储能充放电轨迹，下一6小时严格执行；消除了情景提前获知未来和优化、执行不一致。该方法仍是滚动近似，不宣称严格多阶段随机最优。","","主结果采用：原计划费用保留，每次提交均相对当前合同逐笔收取上调1.5p、下调0.5p的费用。题面对此不够明确，标记为**需人工确认**。","","| 指标（2025-02-01至12-31） | 主策略 |","|---|---:|",f"| 非紧急结算费用/元 | {fmt(t['settlement_cost_yuan'])} |",f"| 紧急购电量/kWh | {fmt(t['emergency_kwh'])} |",f"| 紧急购电费用/元 | {fmt(t['emergency_cost_yuan'])} |",f"| 总费用/元 | {fmt(t['total_cost_yuan'])} |","","## 2. 8种预报组合","","| 日内新预报 | 总费用/元 | 相对仅0:00节省/元 |","|---|---:|---:|"]
    for x in ("none","6","12","18","6+12","6+18","12+18","6+12+18"): L.append(f"| {x} | {fmt(c[x]['total_cost_yuan'])} | {fmt(base-c[x]['total_cost_yuan'])} |")
    L += ["",f"冻结0:00光伏预报但仍在6/12/18时重优化：**{fmt(c['6+12+18_frozen0']['total_cost_yuan'])} 元**。", "","## 3. 结算敏感性","",f"退款/最终净额口径重新优化：**{fmt(c['6+12+18_refund']['total_cost_yuan'])} 元**；主策略固定电量仅改口径重算：**{fmt(s['fixed_quantity_refund_resettlement_total_cost_yuan'])} 元**。均不替代主口径。","","## 4. 指定日期结果",""]
    wanted=[("10:00-10:10",59),("12:00-12:10",71),("14:00-14:10",83),("16:00-16:10",95),("18:00-18:10",107),("20:00-20:10",119)]
    for d in s["target_dates"].values(): L += [f"### {d['date']}","","| 时段 | 计划/kWh | 最终合同/kWh |","|---|---:|---:|"]+[f"| {x} | {fmt(d['plan_kwh'][i])} | {fmt(d['adjusted_kwh'][i])} |" for x,i in wanted]+["","| 4小时段 | 充电/kWh | 放电/kWh |","|---|---:|---:|"]+[f"| {b['time_range']} | {fmt(b['charge_kwh'])} | {fmt(b['discharge_kwh'])} |" for b in d["storage_blocks"]]+["",f"紧急购电合计：{fmt(d['emergency_total_kwh'])} kWh。",""]
    q=s["checks"];L += ["## 5. 校验与限制","",f"最大能量平衡残差 {q['max_energy_balance_residual_kwh']:.3e} kWh，最大SOC递推残差 {q['max_soc_residual_kwh']:.3e} kWh；SOC范围 {fmt(q['soc_min_observed_kwh'])}—{fmt(q['soc_max_observed_kwh'])} kWh；同时充放电 {q['simultaneous_charge_discharge_intervals']} 个时段。","","紧急购电同时充电可能出现，这是共同储能轨迹在实际情景下严格执行的结果。零点计划未用完整情景树联合定价未来调整机会，故只称滚动策略。午夜购电锁定为前一日最后合同量、10分钟光伏线性插值及结算方式均为显式假设。",""]
    return "\n".join(L)
def main():
    p=argparse.ArgumentParser();p.add_argument("--data-dir",type=Path,required=True);p.add_argument("--output-dir",type=Path,required=True);p.add_argument("--report",type=Path,required=True);a=p.parse_args();cfg=Q3Config();a.output_dir.mkdir(parents=True,exist_ok=True);a.report.parent.mkdir(parents=True,exist_ok=True);inputs=load_q3_inputs(a.data_dir.resolve(),cfg);policies={};releases=[]
    for n in range(4):
        for u in itertools.combinations((1,2,3),n): r,z=simulate_policy(inputs,cfg,u);policies[policy_label(u)]=r;releases+=z
    frozen,z=simulate_policy(inputs,cfg,(1,2,3),"frozen0");policies["6+12+18_frozen0"]=frozen;releases+=z
    refund,z=simulate_policy(inputs,cfg,(1,2,3),settlement="final_refund");policies["6+12+18_refund"]=refund;releases+=z
    summary,payload,tables=assemble(policies["6+12+18"],policies,releases,inputs,cfg);(a.output_dir/"summary.json").write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding="utf-8");(a.output_dir/"solver_payload.json").write_text(json.dumps(payload,ensure_ascii=False),encoding="utf-8");(a.output_dir/"release_log.json").write_text(json.dumps(releases,ensure_ascii=False),encoding="utf-8");write_csv(a.output_dir/"daily_metrics.csv",tables["daily"]);write_csv(a.output_dir/"interval_detail.csv",tables["interval"])
    start=inputs["dates"].index(DELIVERY_START);comparison=[]
    for label,rs in policies.items():
        for r in rs[start:]: comparison.append({"policy":label,"date":r["date"].isoformat(),"settlement_cost_yuan":r["settlement_cost"],"emergency_kwh":float(r["emergency"].sum()),"emergency_cost_yuan":r["emergency_cost"],"total_cost_yuan":r["settlement_cost"]+r["emergency_cost"],"soc_start_kwh":float(r["soc"][0]),"soc_end_kwh":float(r["soc"][-1])})
    write_csv(a.output_dir/"policy_comparison_daily.csv",comparison);write_csv(a.output_dir/"baseline_none_interval_detail.csv",detail_rows(policies["none"],inputs,start));write_csv(a.output_dir/"baseline_frozen0_interval_detail.csv",detail_rows(frozen,inputs,start));a.report.write_text(build_report(summary),encoding="utf-8");print(json.dumps(summary["totals_natural_day"],ensure_ascii=False,indent=2))
if __name__=="__main__": main()
