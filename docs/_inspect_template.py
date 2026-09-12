"""摸清附件5/result3.xlsx 的骨架：工作表、尺寸、合并单元格、每张表的行/列内容。"""
from __future__ import annotations

import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
p = ROOT / "problem" / "data" / "附件5" / "result3.xlsx"
print("file:", p, "exists:", p.exists())

wb = openpyxl.load_workbook(p)
print("sheets:", wb.sheetnames)
for name in wb.sheetnames:
    ws = wb[name]
    print()
    print("=" * 100)
    print(f"[{name}]  dims={ws.dimensions}  max_row={ws.max_row} max_col={ws.max_column}")
    print("merges:", [str(m) for m in ws.merged_cells.ranges][:40])
    print("col widths:", {k: v.width for k, v in list(ws.column_dimensions.items())[:8]})
    print("-" * 100)
    for r in range(1, min(ws.max_row, 30) + 1):
        vals = []
        for c in range(1, min(ws.max_column, 12) + 1):
            v = ws.cell(r, c).value
            if v is None:
                vals.append("·")
            else:
                s = repr(v)
                vals.append(s if len(s) <= 22 else s[:20] + "..")
        tail = ""
        if ws.max_column > 12:
            last = [ws.cell(r, c).value for c in (ws.max_column - 1, ws.max_column)]
            tail = "   ...|" + "|".join("·" if v is None else repr(v)[:24] for v in last)
        print(f"  r{r:>3} | " + " | ".join(vals) + tail)
