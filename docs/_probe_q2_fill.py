"""看 Q2/Q3 已填工作簿的填报口径：全天购电量/购电费列、单位、充放电量表结构。"""
from __future__ import annotations

from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
for rel in ("outputs/q2/result2.xlsx", "outputs/q3/result3.xlsx"):
    p = ROOT / rel
    if not p.exists():
        print(f"[missing] {rel}")
        continue
    wb = openpyxl.load_workbook(p)
    print("=" * 100)
    print(rel, wb.sheetnames)
    for name in wb.sheetnames:
        ws = wb[name]
        print(f"--- [{name}] dims={ws.dimensions}")
        for r in range(1, min(ws.max_row, 9) + 1):
            cells = []
            for c in range(1, ws.max_column + 1):
                v = ws.cell(r, c).value
                if c > 4 and c < ws.max_column - 2:
                    continue
                cells.append("·" if v is None else repr(v)[:26])
            print(f"   r{r} | " + " | ".join(cells))
