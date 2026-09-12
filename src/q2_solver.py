from __future__ import annotations

import argparse
import csv
import hashlib
import json
import warnings
from dataclasses import dataclass
from datetime import date, timedelta
from pathlib import Path
from typing import Any

import numpy as np
from openpyxl import load_workbook
from scipy.optimize import linprog
from scipy.sparse import coo_matrix

from efficiency import DEFAULT_EFFICIENCY


@dataclass(frozen=True)
class Config:
    dt_hours: float = 1 / 6
    intervals_per_day: int = 144
    eta_charge: float = DEFAULT_EFFICIENCY.eta_charge
    eta_discharge: float = DEFAULT_EFFICIENCY.eta_discharge
    power_limit_kw: float = 5000.0
    soc_min_kwh: float = 1200.0
    soc_max_kwh: float = 10800.0
    initial_soc_kwh: float = 6000.0
    emergency_price_multiple: float = 5.0
    weight_candidates: tuple[float, ...] = (0.0, 0.25, 0.5, 0.75, 1.0)
    residual_window_candidates: tuple[int, ...] = (7, 14)
    weight_validation_days: int = 28
    distribution_validation_days: int = 7
    throughput_penalty: float = 1e-7

    @property
    def interval_limit_kwh(self) -> float:
        return self.power_limit_kw * self.dt_hours


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def load_inputs(data_dir: Path, cfg: Config) -> dict[str, Any]:
    price_path = data_dir / "附件1.xlsx"
    actual_path = data_dir / "附件2.xlsx"
    price_sheet = load_workbook(price_path, read_only=True, data_only=True).active
    price_rows = list(price_sheet.iter_rows(min_row=2, max_row=145, min_col=1, max_col=4, values_only=True))
    if len(price_rows) != 144:
        raise ValueError("附件1应有144个10分钟时段")
    raw_price = np.asarray([float(row[1]) for row in price_rows])
    cold_raw = np.asarray([(float(row[2]) - float(row[3])) * cfg.dt_hours for row in price_rows])

    wb = load_workbook(actual_path, read_only=True, data_only=True)
    load_rows = list(wb["小区负载"].iter_rows(min_row=2, max_row=366, min_col=1, max_col=145, values_only=True))
    pv_rows = list(wb["光伏发电实际功率"].iter_rows(min_row=2, max_row=366, min_col=1, max_col=145, values_only=True))
    if len(load_rows) != 365 or len(pv_rows) != 365:
        raise ValueError("附件2应含365天数据")
    dates = [row[0].date() for row in load_rows]
    expected = [date(2025, 1, 1) + timedelta(days=i) for i in range(365)]
    if dates != expected or [row[0].date() for row in pv_rows] != expected:
        raise ValueError("附件2日期缺失、重复或两张表不一致")
    load_raw = np.asarray([[float(v) * cfg.dt_hours for v in row[1:]] for row in load_rows])
    pv_raw = np.asarray([[float(v) * cfg.dt_hours for v in row[1:]] for row in pv_rows])
    raw_net = load_raw - pv_raw

    # 原始行是[当日00:10,...,23:50,次日00:00]。自然日午夜必须跨行取值。
    natural = np.full_like(raw_net, np.nan)
    natural[:, 1:] = raw_net[:, :143]
    natural[1:, 0] = raw_net[:-1, 143]
    price_natural = np.r_[raw_price[-1], raw_price[:-1]]
    cold_natural = np.r_[cold_raw[-1], cold_raw[:-1]]
    feb1 = dates.index(date(2025, 2, 1))
    return {
        "dates": dates,
        "net_actual": natural,
        "price_template": raw_price,
        "price_natural": price_natural,
        "cold_start_net": cold_natural,
        "mapping_audit": {
            "checked_midnights": 364,
            "max_abs_error_kwh": float(np.nanmax(np.abs(natural[1:, 0] - raw_net[:-1, 143]))),
            "feb1_midnight_source_date": dates[feb1 - 1].isoformat(),
            "feb1_midnight_source_column": "0:00+1",
            "feb1_midnight_net_kwh": float(natural[feb1, 0]),
            "jan1_midnight_treatment": "附件2缺失；仅初始化仿真用附件1冷启动值，且不进入残差校准",
        },
        "input_hashes": {"附件1.xlsx": file_sha256(price_path), "附件2.xlsx": file_sha256(actual_path)},
    }


def point_forecast(net: np.ndarray, cold: np.ndarray, target: int, cutoff: int, weekday_weight: float) -> np.ndarray:
    if cutoff <= 0:
        return cold.copy()
    origin = date(2025, 1, 1)
    weekday = (origin + timedelta(days=target)).weekday()
    past = np.arange(cutoff)
    same = [int(i) for i in past if (origin + timedelta(days=int(i))).weekday() == weekday][-4:]
    recent = [int(i) for i in past[-7:]]

    def robust_median(indices: list[int]) -> np.ndarray | None:
        if not indices:
            return None
        with np.errstate(all="ignore"), warnings.catch_warnings():
            warnings.simplefilter("ignore", category=RuntimeWarning)
            value = np.nanmedian(net[indices], axis=0)
        return np.where(np.isnan(value), cold, value)

    weekday_curve, recent_curve = robust_median(same), robust_median(recent)
    if weekday_curve is None:
        return recent_curve if recent_curve is not None else cold.copy()
    if recent_curve is None:
        return weekday_curve
    return weekday_weight * weekday_curve + (1.0 - weekday_weight) * recent_curve


def choose_weight(net: np.ndarray, cold: np.ndarray, day: int, cfg: Config) -> float:
    validation = list(range(max(1, day - cfg.weight_validation_days), day))
    if not validation:
        return 0.5
    scores = []
    for weight in cfg.weight_candidates:
        errors = [np.nanmean(np.abs(net[j] - point_forecast(net, cold, j, j, weight))) for j in validation]
        scores.append(float(np.mean(errors)))
    return cfg.weight_candidates[int(np.argmin(scores))]


def delivery_residual(net: np.ndarray, cold: np.ndarray, day: int, weight: float) -> np.ndarray:
    center = np.r_[point_forecast(net, cold, day, day, weight), point_forecast(net, cold, day + 1, day, weight)[0]]
    return np.r_[net[day], net[day + 1, 0]] - center


def energy_score(samples: np.ndarray, observed: np.ndarray) -> float:
    first = np.mean(np.linalg.norm(samples - observed, axis=1))
    pairwise = np.mean(np.linalg.norm(samples[:, None, :] - samples[None, :, :], axis=2))
    return float(first - 0.5 * pairwise)


def choose_window(residuals: dict[int, np.ndarray], day: int, cfg: Config) -> int:
    last_complete = day - 2
    validation = list(range(max(1, last_complete - cfg.distribution_validation_days + 1), last_complete + 1))
    best, best_score = cfg.residual_window_candidates[0], np.inf
    for window in cfg.residual_window_candidates:
        scores = []
        for target in validation:
            train = [j for j in range(max(1, target - window), target) if j in residuals]
            if train:
                scores.append(energy_score(np.stack([residuals[j] for j in train]), residuals[target]))
        score = float(np.mean(scores)) if scores else np.inf
        if score < best_score:
            best, best_score = window, score
    return best


def scenario_set(net: np.ndarray, cold: np.ndarray, day: int, weight: float, cfg: Config) -> tuple[np.ndarray, np.ndarray, int]:
    center = np.r_[point_forecast(net, cold, day, day, weight), point_forecast(net, cold, day + 1, day, weight)[0]]
    # day-1轮廓的末端是当前00:00，计划制定时尚不可用；最晚使用day-2。
    residuals = {j: delivery_residual(net, cold, j, weight) for j in range(1, max(1, day - 1))}
    window = choose_window(residuals, day, cfg)
    indices = [j for j in range(max(1, day - 1 - window), day - 1) if j in residuals]
    if not indices:
        return center[None, :], center, window
    return center[None, :] + np.stack([residuals[j] for j in indices]), center, window


def solve_saa(scenarios: np.ndarray, price_natural: np.ndarray, committed_grid: float, initial_soc: float, terminal_value: float, cfg: Config) -> dict[str, np.ndarray]:
    scenario_count, horizon = scenarios.shape
    if horizon != 145:
        raise ValueError("情景长度必须为145")
    planned_count, block = 144, 5 * horizon + 1
    variable_count = planned_count + scenario_count * block
    objective = np.zeros(variable_count)
    objective[:planned_count] = np.r_[price_natural[1:], price_natural[0]]
    horizon_price = np.r_[price_natural, price_natural[0]]
    for scenario in range(scenario_count):
        base = planned_count + scenario * block
        objective[base : base + 2 * horizon] = cfg.throughput_penalty / scenario_count
        objective[base + 2 * horizon : base + 3 * horizon] = cfg.emergency_price_multiple * horizon_price / scenario_count
        objective[base + 5 * horizon] = -terminal_value / scenario_count

    rows: list[int] = []
    cols: list[int] = []
    values: list[float] = []
    rhs: list[float] = []
    row = 0
    for scenario in range(scenario_count):
        base = planned_count + scenario * block
        charge0, discharge0 = base, base + horizon
        emergency0, surplus0, soc0 = base + 2 * horizon, base + 3 * horizon, base + 4 * horizon
        for h in range(horizon):
            if h > 0:
                rows.append(row); cols.append(h - 1); values.append(1.0)
                rhs.append(float(scenarios[scenario, h]))
            else:
                rhs.append(float(scenarios[scenario, h] - committed_grid))
            for col, value in ((charge0 + h, -1.0), (discharge0 + h, 1.0), (emergency0 + h, 1.0), (surplus0 + h, -1.0)):
                rows.append(row); cols.append(col); values.append(value)
            row += 1
        for h in range(horizon):
            for col, value in ((soc0 + h + 1, 1.0), (soc0 + h, -1.0), (charge0 + h, -cfg.eta_charge), (discharge0 + h, 1.0 / cfg.eta_discharge)):
                rows.append(row); cols.append(col); values.append(value)
            rhs.append(0.0); row += 1
        rows.append(row); cols.append(soc0); values.append(1.0); rhs.append(initial_soc); row += 1

    matrix = coo_matrix((values, (rows, cols)), shape=(row, variable_count)).tocsr()
    bounds: list[tuple[float | None, float | None]] = [(0.0, None)] * planned_count
    for _ in range(scenario_count):
        bounds += [(0.0, cfg.interval_limit_kwh)] * (2 * horizon)
        bounds += [(0.0, None)] * (2 * horizon)
        bounds += [(cfg.soc_min_kwh, cfg.soc_max_kwh)] * (horizon + 1)
    result = linprog(objective, A_eq=matrix, b_eq=np.asarray(rhs), bounds=bounds, method="highs")
    if not result.success:
        raise RuntimeError(f"随机线性规划失败：{result.message}")
    charge, discharge = [], []
    for scenario in range(scenario_count):
        base = planned_count + scenario * block
        charge.append(result.x[base : base + horizon])
        discharge.append(result.x[base + horizon : base + 2 * horizon])
    return {"grid": result.x[:planned_count], "reference_charge": np.mean(charge, axis=0), "reference_discharge": np.mean(discharge, axis=0)}


def causal_dispatch(actual: np.ndarray, grid: np.ndarray, ref_charge: np.ndarray, ref_discharge: np.ndarray, initial_soc: float, cfg: Config) -> dict[str, np.ndarray]:
    count = len(actual)
    charge = np.zeros(count); discharge = np.zeros(count); emergency = np.zeros(count); surplus = np.zeros(count); soc = np.zeros(count + 1); soc[0] = initial_soc
    for t in range(count):
        c, d = max(float(ref_charge[t]), 0.0), max(float(ref_discharge[t]), 0.0)
        common = min(c, d); c -= common; d -= common
        c = min(c, cfg.interval_limit_kwh, (cfg.soc_max_kwh - soc[t]) / cfg.eta_charge)
        d = min(d, cfg.interval_limit_kwh, (soc[t] - cfg.soc_min_kwh) * cfg.eta_discharge)
        gap = float(actual[t] - (grid[t] + d - c))
        if gap > 0:
            reduction = min(c, gap); c -= reduction; gap -= reduction
            extra = max(0.0, min(gap, cfg.interval_limit_kwh - d, (soc[t] - cfg.soc_min_kwh) * cfg.eta_discharge - d))
            d += extra; gap -= extra; emergency[t] = max(gap, 0.0)
        elif gap < 0:
            remainder = -gap
            reduction = min(d, remainder); d -= reduction; remainder -= reduction
            extra = max(0.0, min(remainder, cfg.interval_limit_kwh - c, (cfg.soc_max_kwh - soc[t]) / cfg.eta_charge - c))
            c += extra; remainder -= extra; surplus[t] = max(remainder, 0.0)
        charge[t], discharge[t] = c, d
        soc[t + 1] = soc[t] + cfg.eta_charge * c - d / cfg.eta_discharge
    return {"charge": charge, "discharge": discharge, "emergency": emergency, "surplus": surplus, "soc": soc}


def fixed_reference_dispatch(actual: np.ndarray, grid: np.ndarray, ref_charge: np.ndarray, ref_discharge: np.ndarray, initial_soc: float, cfg: Config) -> dict[str, np.ndarray]:
    """同一起点下不响应实际偏差的逐日局部反事实。"""
    count = len(actual)
    charge = np.zeros(count); discharge = np.zeros(count); emergency = np.zeros(count); surplus = np.zeros(count); soc = np.zeros(count + 1); soc[0] = initial_soc
    for t in range(count):
        c, d = max(float(ref_charge[t]), 0.0), max(float(ref_discharge[t]), 0.0)
        common = min(c, d); c -= common; d -= common
        c = min(c, cfg.interval_limit_kwh, (cfg.soc_max_kwh - soc[t]) / cfg.eta_charge)
        d = min(d, cfg.interval_limit_kwh, (soc[t] - cfg.soc_min_kwh) * cfg.eta_discharge)
        gap = float(actual[t] - (grid[t] + d - c))
        emergency[t] = max(gap, 0.0); surplus[t] = max(-gap, 0.0)
        charge[t], discharge[t] = c, d
        soc[t + 1] = soc[t] + cfg.eta_charge * c - d / cfg.eta_discharge
    return {"charge": charge, "discharge": discharge, "emergency": emergency, "surplus": surplus, "soc": soc}


def no_storage_plan(scenarios: np.ndarray, cfg: Config) -> np.ndarray:
    plans = []
    for h in range(1, 145):
        sample = scenarios[:, h]
        candidates = np.unique(np.r_[0.0, np.maximum(sample, 0.0)])
        loss = candidates + cfg.emergency_price_multiple * np.mean(np.maximum(sample[:, None] - candidates[None, :], 0.0), axis=0)
        plans.append(float(candidates[int(np.argmin(loss))]))
    return np.asarray(plans)


def clock(index: int) -> str:
    if index == 144:
        return "24:00"
    minutes = index * 10
    return f"{minutes // 60:02d}:{minutes % 60:02d}"


def emergency_segments(values: np.ndarray, tolerance: float = 1e-7) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    start: int | None = None
    for index, value in enumerate(np.r_[values, 0.0]):
        if value > tolerance and start is None:
            start = index
        elif value <= tolerance and start is not None:
            result.append({"time_range": f"{clock(start)}-{clock(index)}", "energy_kwh": float(values[start:index].sum())})
            start = None
    return result


def simulate(inputs: dict[str, Any], cfg: Config) -> list[dict[str, Any]]:
    net = inputs["net_actual"]
    net_sim = net.copy(); net_sim[0, 0] = inputs["cold_start_net"][0]
    price = inputs["price_natural"]
    terminal_value = cfg.eta_discharge * float(np.min(price))
    soc_now = cfg.initial_soc_kwh
    committed = max(float(inputs["cold_start_net"][0]), 0.0)
    no_storage_committed = committed
    records: list[dict[str, Any]] = []
    target_set = {date(2025, 3, 20), date(2025, 6, 21), date(2025, 9, 23), date(2025, 12, 21)}
    for day, current_date in enumerate(inputs["dates"]):
        weight = choose_weight(net, inputs["cold_start_net"], day, cfg)
        scenarios, center, window = scenario_set(net, inputs["cold_start_net"], day, weight, cfg)
        solution = solve_saa(scenarios, price, committed, soc_now, terminal_value, cfg)
        published = solution["grid"]
        natural_grid = np.r_[committed, published[:143]]
        dispatch = causal_dispatch(net_sim[day], natural_grid, solution["reference_charge"][:144], solution["reference_discharge"][:144], soc_now, cfg)
        fixed = fixed_reference_dispatch(net_sim[day], natural_grid, solution["reference_charge"][:144], solution["reference_discharge"][:144], soc_now, cfg)
        no_storage_published = no_storage_plan(scenarios, cfg)
        no_storage_grid = np.r_[no_storage_committed, no_storage_published[:143]]
        no_storage_emergency = np.maximum(net_sim[day] - no_storage_grid, 0.0)
        sensitivity = []
        if current_date in target_set:
            for factor in (0.0, 1.0, 2.0):
                alternate = solve_saa(scenarios, price, committed, soc_now, factor * terminal_value, cfg)
                alternate_grid = np.r_[committed, alternate["grid"][:143]]
                alternate_dispatch = causal_dispatch(net_sim[day], alternate_grid, alternate["reference_charge"][:144], alternate["reference_discharge"][:144], soc_now, cfg)
                sensitivity.append({"terminal_value_factor": factor, "natural_day_cost_yuan": float(price @ alternate_grid + cfg.emergency_price_multiple * price @ alternate_dispatch["emergency"]), "soc_end_kwh": float(alternate_dispatch["soc"][-1])})
        records.append({"date": current_date, "weight": weight, "window": window, "scenario_count": len(scenarios), "forecast": center[:144], "published_grid": published, "natural_grid": natural_grid, **dispatch, "fixed_reference": fixed, "no_storage_grid": no_storage_grid, "no_storage_emergency": no_storage_emergency, "sensitivity": sensitivity})
        soc_now = float(dispatch["soc"][-1]); committed = float(published[-1]); no_storage_committed = float(no_storage_published[-1])
        if (day + 1) % 50 == 0:
            print(f"已完成 {day + 1}/365 天")
    return records


def block_sums(values: np.ndarray) -> list[float]:
    return [float(values[i * 24 : (i + 1) * 24].sum()) for i in range(6)]


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0])); writer.writeheader(); writer.writerows(rows)


def assemble(records: list[dict[str, Any]], inputs: dict[str, Any], cfg: Config) -> tuple[dict[str, Any], dict[str, Any], dict[str, list[dict[str, Any]]]]:
    start = inputs["dates"].index(date(2025, 2, 1)); delivered = records[start:]
    price, template_price = inputs["price_natural"], inputs["price_template"]
    interval_rows: list[dict[str, Any]] = []; daily_rows: list[dict[str, Any]] = []; payload_days: list[dict[str, Any]] = []
    max_balance = max_soc = 0.0
    for idx, record in enumerate(delivered, start=start):
        actual = inputs["net_actual"][idx]
        balance = record["natural_grid"] + record["emergency"] + record["discharge"] - record["charge"] - record["surplus"] - actual
        soc_residual = record["soc"][1:] - record["soc"][:-1] - cfg.eta_charge * record["charge"] + record["discharge"] / cfg.eta_discharge
        max_balance = max(max_balance, float(np.max(np.abs(balance)))); max_soc = max(max_soc, float(np.max(np.abs(soc_residual))))
        charges, discharges = block_sums(record["charge"]), block_sums(record["discharge"])
        storage_blocks = [{"time_range": f"{4*i}:00-{4*(i+1)}:00", "charge_kwh": charges[i], "discharge_kwh": discharges[i]} for i in range(6)]
        day = {"date": record["date"].isoformat(), "template_grid_kwh": [float(x) for x in record["published_grid"]], "template_grid_total_kwh": float(record["published_grid"].sum()), "template_grid_cost_yuan": float(template_price @ record["published_grid"]), "natural_grid_total_kwh": float(record["natural_grid"].sum()), "natural_grid_cost_yuan": float(price @ record["natural_grid"]), "soc_start_kwh": float(record["soc"][0]), "soc_end_kwh": float(record["soc"][-1]), "storage_blocks": storage_blocks, "emergency_segments": emergency_segments(record["emergency"]), "emergency_total_kwh": float(record["emergency"].sum()), "emergency_cost_yuan": float(cfg.emergency_price_multiple * price @ record["emergency"]), "sensitivity": record["sensitivity"]}
        payload_days.append(day)
        daily_rows.append({"date": day["date"], "published_plan_kwh": day["template_grid_total_kwh"], "published_plan_cost_yuan": day["template_grid_cost_yuan"], "natural_plan_kwh": day["natural_grid_total_kwh"], "natural_plan_cost_yuan": day["natural_grid_cost_yuan"], "emergency_kwh": day["emergency_total_kwh"], "emergency_cost_yuan": day["emergency_cost_yuan"], "natural_total_cost_yuan": day["natural_grid_cost_yuan"] + day["emergency_cost_yuan"], "surplus_kwh": float(record["surplus"].sum()), "soc_start_kwh": day["soc_start_kwh"], "soc_end_kwh": day["soc_end_kwh"], "forecast_mae_kwh": float(np.mean(np.abs(actual - record["forecast"]))), "weekday_weight": record["weight"], "residual_window_days": record["window"], "scenario_count": record["scenario_count"]})
        for t in range(144):
            interval_rows.append({"date": day["date"], "time_start": clock(t), "price_yuan_per_kwh": float(price[t]), "forecast_net_kwh": float(record["forecast"][t]), "actual_net_kwh": float(actual[t]), "committed_plan_kwh": float(record["natural_grid"][t]), "charge_kwh": float(record["charge"][t]), "discharge_kwh": float(record["discharge"][t]), "soc_start_kwh": float(record["soc"][t]), "soc_end_kwh": float(record["soc"][t+1]), "emergency_kwh": float(record["emergency"][t]), "surplus_kwh": float(record["surplus"][t])})

    all_emergency = np.concatenate([r["emergency"] for r in delivered]); all_charge = np.concatenate([r["charge"] for r in delivered]); all_discharge = np.concatenate([r["discharge"] for r in delivered]); all_soc = np.concatenate([r["soc"] for r in delivered])
    plan_cost = sum(float(price @ r["natural_grid"]) for r in delivered); emergency_cost = sum(float(cfg.emergency_price_multiple * price @ r["emergency"]) for r in delivered)
    no_storage_plan_cost = sum(float(price @ r["no_storage_grid"]) for r in delivered); no_storage_emergency_cost = sum(float(cfg.emergency_price_multiple * price @ r["no_storage_emergency"]) for r in delivered)
    fixed_emergency_cost = sum(float(cfg.emergency_price_multiple * price @ r["fixed_reference"]["emergency"]) for r in delivered)
    fixed_emergency = np.concatenate([r["fixed_reference"]["emergency"] for r in delivered]); fixed_charge = np.concatenate([r["fixed_reference"]["charge"] for r in delivered])
    targets = {day["date"]: day for day in payload_days if day["date"] in {"2025-03-20", "2025-06-21", "2025-09-23", "2025-12-21"}}
    summary = {
        "model": "rolling_direct_empirical_saa_with_causal_storage_recourse", "period": {"start": "2025-02-01", "end": "2025-12-31", "days": 334},
        "time_convention": "区间起点；模板行是当日00:00发布、覆盖00:10至次日00:00的同一版计划",
        "parameters": {"emergency_price_multiple": 5, "eta_charge": 0.9, "eta_discharge": 0.9, "soc_range_kwh": [1200, 10800], "interval_limit_kwh": cfg.interval_limit_kwh, "terminal_value_yuan_per_internal_kwh": cfg.eta_discharge * float(np.min(price))},
        "totals_natural_day": {"planned_purchase_kwh": float(sum(r["natural_grid"].sum() for r in delivered)), "emergency_purchase_kwh": float(all_emergency.sum()), "surplus_kwh": float(sum(r["surplus"].sum() for r in delivered)), "planned_purchase_cost_yuan": plan_cost, "emergency_purchase_cost_yuan": emergency_cost, "total_cost_yuan": plan_cost + emergency_cost, "emergency_intervals": int(np.count_nonzero(all_emergency > 1e-7))},
        "baseline_no_storage": {"planned_purchase_cost_yuan": no_storage_plan_cost, "emergency_purchase_cost_yuan": no_storage_emergency_cost, "total_cost_yuan": no_storage_plan_cost + no_storage_emergency_cost, "saving_yuan": no_storage_plan_cost + no_storage_emergency_cost - plan_cost - emergency_cost},
        "local_counterfactual_fixed_storage": {"definition": "每天从主策略相同SOC起点执行参考充放电，不响应当日实际偏差；为局部反事实，不是独立连续全年策略", "planned_purchase_cost_yuan": plan_cost, "emergency_purchase_cost_yuan": fixed_emergency_cost, "total_cost_yuan": plan_cost + fixed_emergency_cost, "causal_recourse_saving_yuan": fixed_emergency_cost - emergency_cost, "emergency_while_charging_intervals": int(np.count_nonzero((fixed_emergency > 1e-7) & (fixed_charge > 1e-7)))},
        "forecast_metrics": {"mae_kwh_per_interval": float(np.mean([row["forecast_mae_kwh"] for row in daily_rows])), "selected_weight_counts": {str(w): sum(r["weight"] == w for r in delivered) for w in cfg.weight_candidates}, "selected_window_counts": {str(w): sum(r["window"] == w for r in delivered) for w in cfg.residual_window_candidates}},
        "mapping_audit": inputs["mapping_audit"],
        "checks": {"max_energy_balance_residual_kwh": max_balance, "max_soc_residual_kwh": max_soc, "soc_min_observed_kwh": float(all_soc.min()), "soc_max_observed_kwh": float(all_soc.max()), "max_interval_storage_energy_kwh": float(max(all_charge.max(), all_discharge.max())), "simultaneous_charge_discharge_intervals": int(np.count_nonzero((all_charge > 1e-7) & (all_discharge > 1e-7))), "emergency_while_charging_intervals": int(np.count_nonzero((all_emergency > 1e-7) & (all_charge > 1e-7))), "cross_day_soc_continuity": all(abs(delivered[i]["soc"][-1] - delivered[i+1]["soc"][0]) < 1e-7 for i in range(len(delivered)-1))},
        "target_dates": targets, "input_hashes_sha256": inputs["input_hashes"],
    }
    payload = {"metadata": {"time_convention": summary["time_convention"], "input_hashes_sha256": inputs["input_hashes"]}, "days": payload_days}
    return summary, payload, {"interval": interval_rows, "daily": daily_rows}


def fmt(value: float) -> str:
    return f"{value:,.4f}"


def build_report(summary: dict[str, Any]) -> str:
    totals, baseline, fixed, checks, mapping = summary["totals_natural_day"], summary["baseline_no_storage"], summary["local_counterfactual_fixed_storage"], summary["checks"], summary["mapping_audit"]
    lines = ["# 第二问修正后计算报告", "", "## 1. 修正结论", "", "本次结果已修正午夜跨行映射，并以直接经验情景随机线性规划替代80%分位数等价法。每天0:00只使用已经完成的历史轮廓；实际值仅在对应10分钟时段到达后用于因果储能补救和紧急购电结算。", "", "| 2025-02-01至12-31自然日指标 | 结果 |", "|---|---:|", f"| 计划购电量 | {fmt(totals['planned_purchase_kwh'])} kWh |", f"| 紧急购电量 | {fmt(totals['emergency_purchase_kwh'])} kWh |", f"| 计划购电费 | {fmt(totals['planned_purchase_cost_yuan'])} 元 |", f"| 紧急购电费 | {fmt(totals['emergency_purchase_cost_yuan'])} 元 |", f"| 总购电费 | {fmt(totals['total_cost_yuan'])} 元 |", "", "## 2. 原始时标核验", "", f"共核验{mapping['checked_midnights']}个可核对午夜，最大映射误差为{mapping['max_abs_error_kwh']:.3e} kWh。2025-02-01 00:00取自2025-01-31原始行的`0:00+1`列，实际净负荷为{mapping['feb1_midnight_net_kwh']:.4f} kWh。2025-01-01 00:00在附件2缺失，只用附件1值初始化，且不进入残差样本。", "", "## 3. 模型与对照", "", "主模型保留每条历史145时段残差轮廓，直接优化计划购电、情景充放电和情景紧急购电。日内执行采用因果补救：缺电时先取消充电、再增加可行放电，最后才紧急购电；富余时先取消放电、再增加可行充电。", "", "| 策略 | 总费用/元 |", "|---|---:|", f"| 直接情景优化+因果储能补救 | {fmt(totals['total_cost_yuan'])} |", f"| 固定参考储能的逐日局部反事实 | {fmt(fixed['total_cost_yuan'])} |", f"| 因果补救相对局部反事实节省 | {fmt(fixed['causal_recourse_saving_yuan'])} |", f"| 无储能直接情景基线 | {fmt(baseline['total_cost_yuan'])} |", f"| 主策略相对无储能节省 | {fmt(baseline['saving_yuan'])} |", "", f"固定参考储能反事实中有{fixed['emergency_while_charging_intervals']}个时段同时紧急购电和充电；主策略为0。该反事实每天使用主策略相同SOC起点，只衡量日内补救的局部价值，不冒充独立连续全年策略。", "", "## 4. 指定日期表1：计划购电", "", "以下每行均为该日0:00发布的同一版计划；全天合计覆盖模板的00:10至次日00:00。", ""]
    wanted = [("10:00-10:10", 59), ("12:00-12:10", 71), ("14:00-14:10", 83), ("16:00-16:10", 95), ("18:00-18:10", 107), ("20:00-20:10", 119)]
    for day in summary["target_dates"].values():
        lines += [f"### {day['date']}", "", "| 时间段 | 购电量/kWh |", "|---|---:|"] + [f"| {label} | {fmt(day['template_grid_kwh'][index])} |" for label, index in wanted] + [f"| 全天购电量 | {fmt(day['template_grid_total_kwh'])} |", f"| 全天购电费 | {fmt(day['template_grid_cost_yuan'])} 元 |", ""]
    lines += ["## 5. 指定日期表2：实际充放电", ""]
    for day in summary["target_dates"].values():
        lines += [f"### {day['date']}", "", "| 时间段 | 充电量/kWh | 放电量/kWh |", "|---|---:|---:|"] + [f"| {item['time_range']} | {fmt(item['charge_kwh'])} | {fmt(item['discharge_kwh'])} |" for item in day["storage_blocks"]] + [f"| 0:00储电量 | {fmt(day['soc_start_kwh'])} | |", f"| 24:00储电量 | {fmt(day['soc_end_kwh'])} | |", ""]
    lines += ["## 6. 指定日期表3：紧急购电", ""]
    for day in summary["target_dates"].values():
        lines += [f"### {day['date']}", ""]
        if day["emergency_segments"]:
            lines += ["| 时间段 | 购电量/kWh |", "|---|---:|"] + [f"| {item['time_range']} | {fmt(item['energy_kwh'])} |" for item in day["emergency_segments"]]
        else:
            lines += ["无紧急购电。"]
        lines.append("")
    lines += ["## 7. 终端价值敏感性", "", "已删除48小时强制复位。主模型采用保守终端价值`0.9×最低电价`，并在四个指定日期以相同初始状态比较0、1、2倍终端价值。", "", "| 日期 | 倍数 | 当日费用/元 | 24:00 SOC/kWh |", "|---|---:|---:|---:|"]
    for day in summary["target_dates"].values():
        for item in day["sensitivity"]:
            lines.append(f"| {day['date']} | {item['terminal_value_factor']:.0f} | {fmt(item['natural_day_cost_yuan'])} | {fmt(item['soc_end_kwh'])} |")
    lines += ["", "## 8. 数值校验", "", "| 校验项 | 结果 |", "|---|---:|", f"| 最大能量平衡残差 | {checks['max_energy_balance_residual_kwh']:.3e} kWh |", f"| 最大SOC递推残差 | {checks['max_soc_residual_kwh']:.3e} kWh |", f"| SOC范围 | {fmt(checks['soc_min_observed_kwh'])}—{fmt(checks['soc_max_observed_kwh'])} kWh |", f"| 最大单时段充放电量 | {fmt(checks['max_interval_storage_energy_kwh'])} kWh |", f"| 同时充放电时段 | {checks['simultaneous_charge_discharge_intervals']} |", f"| 紧急购电同时充电时段 | {checks['emergency_while_charging_intervals']} |", f"| 跨日SOC连续 | {'是' if checks['cross_day_soc_continuity'] else '否'} |", "", "## 9. 限制", "", "附件未提供0:00可用的负荷预测，因此预测器、历史窗口和终端价值均是可复现的建模假设。滚动验证只用于选择候选权重和残差窗口，不把逐时段覆盖率解释为全天可靠性。终端敏感性仅覆盖四个指定日期，不能替代更大范围的鲁棒性研究。", ""]
    return "\n".join(lines)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="第二问直接经验情景随机优化")
    parser.add_argument("--data-dir", type=Path, required=True); parser.add_argument("--output-dir", type=Path, required=True); parser.add_argument("--report", type=Path, required=True)
    return parser.parse_args()


def main() -> None:
    args = parse_args(); cfg = Config(); args.output_dir.mkdir(parents=True, exist_ok=True); args.report.parent.mkdir(parents=True, exist_ok=True)
    inputs = load_inputs(args.data_dir.resolve(), cfg); records = simulate(inputs, cfg); summary, payload, tables = assemble(records, inputs, cfg)
    (args.output_dir / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    (args.output_dir / "solver_payload.json").write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    write_csv(args.output_dir / "daily_metrics.csv", tables["daily"]); write_csv(args.output_dir / "interval_detail.csv", tables["interval"])
    args.report.write_text(build_report(summary), encoding="utf-8")
    print(json.dumps(summary["totals_natural_day"], ensure_ascii=False, indent=2)); print(f"报告：{args.report.resolve()}")


if __name__ == "__main__":
    main()
