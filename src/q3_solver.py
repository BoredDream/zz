from __future__ import annotations

import argparse
import csv
import itertools
import json
from dataclasses import asdict, dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

import numpy as np
from openpyxl import load_workbook
from scipy.optimize import linprog
from scipy.sparse import coo_matrix

from q2_solver import Config, emergency_segments, energy_score, file_sha256, point_forecast

MODEL_VERSION = "q3_causal_rolling_saa_common_storage_v3"
ISSUE_INTERVALS = (0, 36, 72, 108)
ISSUE_LABELS = ("0:00", "6:00", "12:00", "18:00")
TARGET_DATES = {date(2025, 3, 20), date(2025, 6, 21), date(2025, 9, 23), date(2025, 12, 21)}
DELIVERY_START = date(2025, 2, 1)

WAIT_PERIODS = (("10:00-10:10", 59), ("12:00-12:10", 71), ("14:00-14:10", 83), ("16:00-16:10", 95), ("18:00-18:10", 107), ("20:00-20:10", 119))


@dataclass(frozen=True)
class Q3Config(Config):
    pv_interpolation: str = "linear_from_latest_observation_to_hourly_endpoints"


def parse_date(value: Any) -> date:
    return value.date() if isinstance(value, datetime) else datetime.strptime(str(value), "%Y-%m-%d").date()


def load_q3_inputs(data_dir: Path, cfg: Q3Config) -> dict[str, Any]:
    price_path, actual_path, forecast_path = data_dir / "附件1.xlsx", data_dir / "附件2.xlsx", data_dir / "附件3.xlsx"
    price_rows = list(load_workbook(price_path, read_only=True, data_only=True).active.iter_rows(min_row=2, max_row=145, min_col=1, max_col=4, values_only=True))
    raw_price = np.asarray([float(row[1]) for row in price_rows])
    cold_load_raw = np.asarray([float(row[2]) for row in price_rows])
    cold_pv_raw = np.asarray([float(row[3]) for row in price_rows])
    workbook = load_workbook(actual_path, read_only=True, data_only=True)
    load_rows = list(workbook["小区负载"].iter_rows(min_row=2, max_row=366, min_col=1, max_col=145, values_only=True))
    pv_rows = list(workbook["光伏发电实际功率"].iter_rows(min_row=2, max_row=366, min_col=1, max_col=145, values_only=True))
    dates = [row[0].date() for row in load_rows]
    expected = [date(2025, 1, 1) + timedelta(days=i) for i in range(365)]
    if dates != expected or [row[0].date() for row in pv_rows] != expected:
        raise ValueError("附件2日期不连续或两表日期不一致")
    load_raw = np.asarray([[float(v) for v in row[1:]] for row in load_rows])
    pv_raw = np.asarray([[float(v) for v in row[1:]] for row in pv_rows])
    # 原始行是[0:10,...,23:50,次日0:00]。自然日00:00必须取上一原始行末列，禁止同行循环移位。
    load_natural = np.full_like(load_raw, np.nan)
    pv_natural = np.full_like(pv_raw, np.nan)
    load_natural[:, 1:] = load_raw[:, :143]
    pv_natural[:, 1:] = pv_raw[:, :143]
    load_natural[1:, 0] = load_raw[:-1, 143]
    pv_natural[1:, 0] = pv_raw[:-1, 143]
    cold_load = np.r_[cold_load_raw[-1], cold_load_raw[:-1]]
    cold_pv = np.r_[cold_pv_raw[-1], cold_pv_raw[:-1]]
    rows = list(load_workbook(forecast_path, read_only=True, data_only=True).active.iter_rows(min_row=2, max_row=1461, min_col=1, max_col=26, values_only=True))
    if len(rows) != 1460:
        raise ValueError("附件3应有365天×4个发布时刻")
    forecasts = np.zeros((365, 4, 24))
    forecast_dates: list[date] = []
    for day in range(365):
        group = rows[4 * day : 4 * day + 4]
        forecast_dates.append(parse_date(group[0][0]))
        if [str(row[1]) for row in group] != list(ISSUE_LABELS):
            raise ValueError(f"附件3第{day + 1}天发布时刻不完整")
        forecasts[day] = np.asarray([[float(v) for v in row[2:26]] for row in group])
    if forecast_dates != expected:
        raise ValueError("附件3日期不连续")
    feb = dates.index(DELIVERY_START)
    return {
        "dates": dates,
        "price_natural": np.r_[raw_price[-1], raw_price[:-1]],
        "price_template": raw_price,
        "load_actual_kw": load_natural,
        "pv_actual_kw": pv_natural,
        "net_actual_kwh": (load_natural - pv_natural) * cfg.dt_hours,
        "cold_load_kw": cold_load,
        "cold_pv_kw": cold_pv,
        "pv_forecast_kw": forecasts,
        "mapping_audit": {
            "actual_midnights_checked": 364,
            "actual_midnight_max_error_kw": float(max(np.nanmax(abs(load_natural[1:, 0] - load_raw[:-1, 143])), np.nanmax(abs(pv_natural[1:, 0] - pv_raw[:-1, 143])))),
            "feb1_midnight_source": "附件2 2025-01-31 行末列(0:00+1)",
            "feb1_midnight_net_kwh": float((load_natural[feb, 0] - pv_natural[feb, 0]) * cfg.dt_hours),
            "natural_day": "00:00—24:00，共144个10分钟区间",
            "template_row": "00:10—次日00:00，共144个10分钟区间；末列次日00:00等于下一自然日00:00",
            "template_row_coverage": "模板行覆盖自然日00:10至次日00:10，不与自然日逐项相同，两者费用不能跨表相加",
        },
        "input_hashes": {p.name: file_sha256(p) for p in (price_path, actual_path, forecast_path)},
    }


def absolute_profile(values: np.ndarray, cold: np.ndarray, day: int, weight: float) -> np.ndarray:
    """自然日00:00—次日00:00的145个节点负荷预测（仅使用已完整出现的历史日）。"""
    return np.r_[point_forecast(values, cold, day, day, weight), point_forecast(values, cold, day + 1, day, weight)[0]]


def pv_curve(inputs: dict[str, Any], day: int, source_i: int) -> np.ndarray:
    """发布时刻source_i起，自然日issue..24:00的145-source个节点光伏预报。"""
    issue = ISSUE_INTERVALS[source_i]
    if issue == 0:
        anchor = float(inputs["pv_actual_kw"][day - 1, 143]) if day > 0 else float(inputs["cold_pv_kw"][0])
    else:
        anchor = float(inputs["pv_actual_kw"][day, issue - 1])
    return np.interp(np.arange(145 - issue), np.arange(0, 145, 6), np.r_[anchor, inputs["pv_forecast_kw"][day, source_i]])


class ForecastCache:
    """与策略无关的预测与情景缓存；10套策略共用，避免重复计算。"""

    def __init__(self, inputs: dict[str, Any], cfg: Q3Config):
        self.inputs, self.cfg = inputs, cfg
        self._weight: dict[int, float] = {}
        self._absolute: dict[tuple[int, float], np.ndarray] = {}
        self._pv: dict[tuple[int, int], np.ndarray] = {}
        self._scenarios: dict[tuple[int, int, int], tuple[np.ndarray, np.ndarray, int]] = {}

    def weight(self, day: int) -> float:
        if day not in self._weight:
            validation = list(range(max(1, day - self.cfg.weight_validation_days), day))
            if not validation:
                self._weight[day] = 0.5
            else:
                scores = [float(np.mean([np.nanmean(np.abs(self.inputs["load_actual_kw"][j] - point_forecast(self.inputs["load_actual_kw"], self.inputs["cold_load_kw"], j, j, w))) for j in validation])) for w in self.cfg.weight_candidates]
                self._weight[day] = self.cfg.weight_candidates[int(np.argmin(scores))]
        return self._weight[day]

    def absolute(self, day: int, weight: float) -> np.ndarray:
        key = (day, weight)
        if key not in self._absolute:
            self._absolute[key] = absolute_profile(self.inputs["load_actual_kw"], self.inputs["cold_load_kw"], day, weight)
        return self._absolute[key]

    def pv(self, day: int, source_i: int) -> np.ndarray:
        key = (day, source_i)
        if key not in self._pv:
            self._pv[key] = pv_curve(self.inputs, day, source_i)
        return self._pv[key]

    def net_forecast(self, day: int, target_i: int, source_i: int, weight: float) -> np.ndarray:
        target, source = ISSUE_INTERVALS[target_i], ISSUE_INTERVALS[source_i]
        if source > target:
            raise ValueError("预报源时刻晚于决策时刻")
        return (self.absolute(day, weight)[target:] - self.pv(day, source_i)[target - source:]) * self.cfg.dt_hours

    def actual_horizon(self, day: int, target_i: int) -> np.ndarray:
        issue = ISSUE_INTERVALS[target_i]
        return np.r_[self.inputs["net_actual_kwh"][day, issue:], self.inputs["net_actual_kwh"][day + 1, 0]]

    def scenarios(self, day: int, target_i: int, source_i: int) -> tuple[np.ndarray, np.ndarray, int]:
        key = (day, target_i, source_i)
        if key not in self._scenarios:
            weight = self.weight(day)
            center = self.net_forecast(day, target_i, source_i, weight)
            latest = day - 2 if target_i == 0 else day - 1
            # 残差轮廓必须与中心预测使用同一个负荷预测器（同一权重），否则情景分布与实际预测误差不一致。
            residuals = {}
            for j in range(1, max(1, latest + 1)):
                residuals[j] = self.actual_horizon(j, target_i) - self.net_forecast(j, target_i, source_i, weight)
            window = self._choose_window(residuals, latest)
            selected = [j for j in range(max(1, latest - window + 1), latest + 1) if j in residuals]
            matrix = center[None, :] if not selected else center[None, :] + np.stack([residuals[j] for j in selected])
            self._scenarios[key] = (matrix, center, window)
        return self._scenarios[key]

    def _choose_window(self, residuals: dict[int, np.ndarray], latest: int) -> int:
        best, best_score = self.cfg.residual_window_candidates[0], np.inf
        for window in self.cfg.residual_window_candidates:
            scores = []
            for target in range(max(1, latest - self.cfg.distribution_validation_days + 1), latest + 1):
                train = [j for j in range(max(1, target - window), target) if j in residuals]
                if train:
                    scores.append(energy_score(np.stack([residuals[j] for j in train]), residuals[target]))
            score = float(np.mean(scores)) if scores else np.inf
            if score < best_score:
                best, best_score = window, score
        return best


def solve_common(scenarios: np.ndarray, prices: np.ndarray, initial_soc: float, terminal: float, cfg: Q3Config, fixed_first: float | None = None, reference: np.ndarray | None = None, mode: str = "plan") -> dict[str, np.ndarray]:
    """同一发布时刻的全部情景共享购电、充电、放电和SOC；仅紧急购电与弃电随情景变化。

    索引约定：h=0 对应锁定区间，h>=1 对应计划向量第 h-1 项。
    计划模式的节点编号为自然日区间 0..H-1；调整模式为自然日区间 issue..issue+H-1。
    末时段(index H-1，即次日00:00—00:10)的储能动作由次日模型决定，本日不提交，固定为0。
    """
    width, horizon = scenarios.shape
    grid_count = horizon - 1 if mode == "plan" else horizon
    increment0 = grid_count
    decrement0 = increment0 + (0 if mode == "plan" else grid_count)
    common = grid_count if mode == "plan" else 3 * grid_count
    charge0, discharge0 = common, common + horizon
    soc0 = discharge0 + horizon
    recourse = soc0 + horizon + 1
    size = recourse + 2 * width * horizon
    objective = np.zeros(size)
    if mode == "plan":
        objective[:grid_count] = prices[1:]
    elif mode == "per_submission":
        objective[increment0 : increment0 + grid_count] = 1.5 * prices
        objective[decrement0 : decrement0 + grid_count] = 0.5 * prices
    elif mode == "final_refund":
        objective[:grid_count] = prices
        objective[increment0 : increment0 + grid_count] = 0.5 * prices
        objective[decrement0 : decrement0 + grid_count] = 0.5 * prices
    else:
        raise ValueError(mode)
    objective[charge0 : charge0 + horizon] = cfg.throughput_penalty
    objective[discharge0 : discharge0 + horizon] = cfg.throughput_penalty
    objective[soc0 + horizon - 1] = -terminal
    for w in range(width):
        objective[recourse + 2 * w * horizon : recourse + (2 * w + 1) * horizon] = cfg.emergency_price_multiple * prices / width

    rows: list[int] = []
    cols: list[int] = []
    values: list[float] = []
    rhs: list[float] = []
    row = 0
    if mode != "plan":
        if reference is None or len(reference) != grid_count:
            raise ValueError("reference长度错误")
        for h in range(grid_count):
            for col, value in ((h, 1.0), (increment0 + h, -1.0), (decrement0 + h, 1.0)):
                rows.append(row); cols.append(col); values.append(value)
            rhs.append(float(reference[h])); row += 1
    for w in range(width):
        emergency0, surplus0 = recourse + 2 * w * horizon, recourse + (2 * w + 1) * horizon
        for h in range(horizon):
            if mode == "plan":
                if h == 0:
                    const = float(fixed_first)
                else:
                    rows.append(row); cols.append(h - 1); values.append(1.0); const = 0.0
            else:
                rows.append(row); cols.append(h); values.append(1.0); const = 0.0
            for col, value in ((charge0 + h, -1.0), (discharge0 + h, 1.0), (emergency0 + h, 1.0), (surplus0 + h, -1.0)):
                rows.append(row); cols.append(col); values.append(value)
            rhs.append(float(scenarios[w, h] - const)); row += 1
    for h in range(horizon):
        for col, value in ((soc0 + h + 1, 1.0), (soc0 + h, -1.0), (charge0 + h, -cfg.eta_charge), (discharge0 + h, 1.0 / cfg.eta_discharge)):
            rows.append(row); cols.append(col); values.append(value)
        rhs.append(0.0); row += 1
    rows.append(row); cols.append(soc0); values.append(1.0); rhs.append(initial_soc); row += 1

    matrix = coo_matrix((values, (rows, cols)), shape=(row, size)).tocsr()
    storage_bounds = [(0.0, cfg.interval_limit_kwh)] * horizon
    storage_bounds[-1] = (0.0, 0.0)
    bounds: list[tuple[float, float]] = [(0.0, None)] * grid_count
    if mode != "plan":
        bounds += [(0.0, None)] * (2 * grid_count)
    bounds += list(storage_bounds) + list(storage_bounds)
    bounds += [(cfg.soc_min_kwh, cfg.soc_max_kwh)] * (horizon + 1)
    bounds += [(0.0, None)] * (2 * width * horizon)
    result = linprog(objective, A_eq=matrix, b_eq=np.asarray(rhs), bounds=bounds, method="highs")
    if not result.success:
        raise RuntimeError(f"随机线性规划失败：{result.message}")
    out = {"grid": result.x[:grid_count], "charge": result.x[charge0 : charge0 + horizon], "discharge": result.x[discharge0 : discharge0 + horizon], "soc": result.x[soc0 : soc0 + horizon + 1], "emergency": np.asarray([result.x[recourse + 2 * w * horizon : recourse + (2 * w + 1) * horizon] for w in range(width)]), "surplus": np.asarray([result.x[recourse + (2 * w + 1) * horizon : recourse + (2 * w + 2) * horizon] for w in range(width)]), "objective": float(result.fun), "status": ("optimal" if result.success else "failed")}
    if mode != "plan":
        out["increase"] = result.x[increment0 : increment0 + grid_count]
        out["decrease"] = result.x[decrement0 : decrement0 + grid_count]
    return out


def execute_fixed(actual: np.ndarray, grid: np.ndarray, charge: np.ndarray, discharge: np.ndarray, soc0: float, cfg: Q3Config) -> dict[str, np.ndarray]:
    """严格执行发布时锁定的共同储能轨迹；实际净负荷到达后只计算紧急购电或弃电。"""
    soc = np.zeros(len(actual) + 1)
    soc[0] = soc0
    for t in range(len(actual)):
        soc[t + 1] = soc[t] + cfg.eta_charge * charge[t] - discharge[t] / cfg.eta_discharge
    gap = actual - (grid + discharge - charge)
    return {"charge": charge.copy(), "discharge": discharge.copy(), "emergency": np.maximum(gap, 0.0), "surplus": np.maximum(-gap, 0.0), "soc": soc}


def policy_label(updates: tuple[int, ...], forecast_mode: str = "latest", settlement: str = "per_submission") -> str:
    label = "none" if not updates else "+".join(str(ISSUE_INTERVALS[i] // 6) for i in updates)
    if forecast_mode == "frozen0":
        label += "_frozen0"
    if settlement == "final_refund":
        label += "_refund"
    return label


def policy_specs() -> list[tuple[str, tuple[int, ...], str, str]]:
    specs = [(policy_label(u), u, "latest", "per_submission") for n in range(4) for u in itertools.combinations((1, 2, 3), n)]
    specs.append(("6+12+18_frozen0", (1, 2, 3), "frozen0", "per_submission"))
    specs.append(("6+12+18_refund", (1, 2, 3), "latest", "final_refund"))
    return specs


def simulate_policy(inputs: dict[str, Any], cfg: Q3Config, updates: tuple[int, ...], forecast_mode: str, settlement: str, cache: ForecastCache) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    label = policy_label(updates, forecast_mode, settlement)
    net = inputs["net_actual_kwh"].copy()
    net[0, 0] = (inputs["cold_load_kw"][0] - inputs["cold_pv_kw"][0]) * cfg.dt_hours
    price = inputs["price_natural"]
    terminal = cfg.eta_discharge * float(price.min())
    soc_now = cfg.initial_soc_kwh
    prev_plan = max(float(net[0, 0]), 0.0)
    prev_active, prev_penalty = prev_plan, 0.0
    records: list[dict[str, Any]] = []
    releases: list[dict[str, Any]] = []
    for day, current in enumerate(inputs["dates"]):
        weight = cache.weight(day)
        scenarios0, center0, window0 = cache.scenarios(day, 0, 0)
        plan_solution = solve_common(scenarios0, np.r_[price, price[0]], soc_now, terminal, cfg, fixed_first=prev_active)
        plan = plan_solution["grid"].copy()
        active = plan.copy()
        penalty = np.zeros(144)
        charge = plan_solution["charge"].copy()
        discharge = plan_solution["discharge"].copy()
        trace: dict[str, list[float]] = {key: [] for key in ("charge", "discharge", "emergency", "surplus")}
        trace["soc"] = [soc_now]
        grids: list[float] = []
        metrics = [(0, center0[:36], net[day, :36])]
        windows = {"0:00": window0}
        counts = {"0:00": len(scenarios0)}
        releases.append({"policy": label, "date": current.isoformat(), "issue": "0:00", "forecast_source": "0:00", "reference_kind": "new_plan", "information_cutoff": "previous_day_23:50", "latest_observed_pv_interval": "previous_day_23:50_24:00", "history_cutoff": "day-2_complete_residuals", "window_days": window0, "scenario_count": len(scenarios0), "solve_status": "optimal", "increase_kwh": 0.0, "decrease_kwh": 0.0, "adjustment_fee_yuan": 0.0, "contract_kwh": list(map(float, plan)), "common_charge_kwh": list(map(float, charge)), "common_discharge_kwh": list(map(float, discharge))})
        for issue_i, issue in enumerate(ISSUE_INTERVALS):
            if issue_i > 0 and issue_i in updates:
                source = issue_i if forecast_mode == "latest" else 0
                scenarios, center, window = cache.scenarios(day, issue_i, source)
                prices = np.r_[price[issue:], price[0]]
                previous = active[issue - 1 :].copy()
                reference = previous if settlement == "per_submission" else plan[issue - 1 :].copy()
                solution = solve_common(scenarios, prices, soc_now, terminal, cfg, reference=reference, mode=settlement)
                active[issue - 1 :] = solution["grid"]
                charge[issue:] = solution["charge"]
                discharge[issue:] = solution["discharge"]
                increase = np.maximum(solution["grid"] - previous, 0.0)
                decrease = np.maximum(previous - solution["grid"], 0.0)
                fee = 1.5 * prices * increase + 0.5 * prices * decrease if settlement == "per_submission" else np.zeros_like(increase)
                if settlement == "per_submission":
                    penalty[issue - 1 :] += fee
                releases.append({"policy": label, "date": current.isoformat(), "issue": ISSUE_LABELS[issue_i], "forecast_source": ISSUE_LABELS[source], "reference_kind": "previous_active" if settlement == "per_submission" else "original_plan", "information_cutoff": f"{issue // 6 - 1:02d}:50", "latest_observed_pv_interval": f"{issue // 6 - 1:02d}:50-{issue // 6:02d}:00", "history_cutoff": "day-1_complete_residuals", "window_days": window, "scenario_count": len(scenarios), "solve_status": "optimal", "increase_kwh": float(increase.sum()), "decrease_kwh": float(decrease.sum()), "adjustment_fee_yuan": float(fee.sum()), "contract_kwh": list(map(float, solution["grid"])), "common_charge_kwh": list(map(float, solution["charge"])), "common_discharge_kwh": list(map(float, solution["discharge"]))})
                metrics.append((issue_i, center[: min(36, 144 - issue)], net[day, issue : issue + 36]))
                windows[ISSUE_LABELS[issue_i]] = window
                counts[ISSUE_LABELS[issue_i]] = len(scenarios)
            block = min(36, 144 - issue)
            grid_block = np.r_[prev_active, active[:35]] if issue == 0 else active[issue - 1 : issue - 1 + block]
            run = execute_fixed(net[day, issue : issue + block], grid_block, charge[issue : issue + block], discharge[issue : issue + block], soc_now, cfg)
            grids.extend(map(float, grid_block))
            for key in ("charge", "discharge", "emergency", "surplus"):
                trace[key].extend(map(float, run[key]))
            trace["soc"].extend(map(float, run["soc"][1:]))
            soc_now = float(run["soc"][-1])
        base = np.r_[prev_plan, plan[:143]]
        adjusted = np.r_[prev_active, active[:143]]
        pen_natural = np.r_[prev_penalty, penalty[:143]]
        emergency = np.asarray(trace["emergency"])
        emergency_cost = float(cfg.emergency_price_multiple * price @ emergency)
        settlement_cost = float(price @ base + pen_natural.sum()) if settlement == "per_submission" else float(price @ (adjusted + 0.5 * np.abs(adjusted - base)))
        records.append({"date": current, "policy": label, "weight": weight, "windows": windows, "scenario_counts": counts, "plan": plan, "adjusted": active, "base_natural": base, "adjusted_natural": adjusted, "penalty_natural": pen_natural, "penalty_template": penalty.copy(), "grid": np.asarray(grids), "charge": np.asarray(trace["charge"]), "discharge": np.asarray(trace["discharge"]), "emergency": emergency, "surplus": np.asarray(trace["surplus"]), "soc": np.asarray(trace["soc"]), "settlement_cost": settlement_cost, "emergency_cost": emergency_cost, "issue_metrics": metrics})
        prev_plan, prev_active, prev_penalty = float(plan[-1]), float(active[-1]), float(penalty[-1])
        if (day + 1) % 100 == 0:
            print(f"  {label}: {day + 1}/365", flush=True)
    return records, releases


RECORD_ARRAYS = ("plan", "adjusted", "base_natural", "adjusted_natural", "penalty_natural", "penalty_template", "grid", "charge", "discharge", "emergency", "surplus", "soc")


def serialize_records(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    payload = []
    for record in records:
        item: dict[str, Any] = {"date": record["date"].isoformat(), "policy": record["policy"], "weight": record["weight"], "windows": record["windows"], "scenario_counts": record["scenario_counts"], "settlement_cost": record["settlement_cost"], "emergency_cost": record["emergency_cost"], "issue_metrics": [[int(ii), list(map(float, f)), list(map(float, o))] for ii, f, o in record["issue_metrics"]]}
        for key in RECORD_ARRAYS:
            item[key] = list(map(float, record[key]))
        payload.append(item)
    return payload


def deserialize_records(payload: list[dict[str, Any]]) -> list[dict[str, Any]]:
    records = []
    for item in payload:
        record = dict(item)
        record["date"] = date.fromisoformat(item["date"])
        for key in RECORD_ARRAYS:
            record[key] = np.asarray(item[key], dtype=float)
        record["issue_metrics"] = [(int(ii), np.asarray(f, dtype=float), np.asarray(o, dtype=float)) for ii, f, o in item["issue_metrics"]]
        records.append(record)
    return records


def build_signature(cfg: Q3Config, inputs: dict[str, Any]) -> str:
    solver = Path(__file__).resolve()
    signature = {
        "model_version": MODEL_VERSION,
        "solver_sha256": file_sha256(solver),
        "q2_solver_sha256": file_sha256(solver.parent / "q2_solver.py"),
        "input_files_sha256": dict(inputs["input_hashes"]),
        "parameters": {key: (list(value) if isinstance(value, tuple) else value) for key, value in asdict(cfg).items()},
        "policies": [[label, list(updates), forecast_mode, settlement] for label, updates, forecast_mode, settlement in policy_specs()],
    }
    return json.dumps(signature, ensure_ascii=False, sort_keys=True)


def load_checkpoint(path: Path, signature: str) -> dict[str, Any] | None:
    if not path.exists():
        return None
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError):
        return None
    if data.get("signature") != signature:
        print(f"  [checkpoint] 签名不一致，忽略 {path.name}")
        return None
    return data


def save_checkpoint(path: Path, signature: str, label: str, records: list[dict[str, Any]], releases: list[dict[str, Any]]) -> None:
    data = {"signature": signature, "policy": label, "model_version": MODEL_VERSION, "days": serialize_records(records), "releases": releases}
    path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")


def block_sums(values: np.ndarray) -> list[float]:
    return [float(values[24 * i : 24 * (i + 1)].sum()) for i in range(6)]


def clock(index: int) -> str:
    return "24:00" if index == 144 else f"{index // 6:02d}:{(index % 6) * 10:02d}"


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0])); writer.writeheader(); writer.writerows(rows)


def totals(records: list[dict[str, Any]], start: int) -> dict[str, float]:
    delivered = records[start:]
    settlement = sum(record["settlement_cost"] for record in delivered)
    emergency = sum(record["emergency_cost"] for record in delivered)
    return {"settlement_cost_yuan": float(settlement), "emergency_cost_yuan": float(emergency), "total_cost_yuan": float(settlement + emergency), "emergency_kwh": float(sum(record["emergency"].sum() for record in delivered)), "planned_purchase_kwh": float(sum(record["base_natural"].sum() for record in delivered)), "final_contract_kwh": float(sum(record["adjusted_natural"].sum() for record in delivered)), "days": len(delivered)}


def detail_rows(records: list[dict[str, Any]], inputs: dict[str, Any], start: int) -> list[dict[str, Any]]:
    rows = []
    price = inputs["price_natural"]
    for index, record in enumerate(records[start:], start=start):
        for t in range(144):
            rows.append({"policy": record["policy"], "date": record["date"].isoformat(), "time_start": clock(t), "price_yuan_per_kwh": float(price[t]), "initial_plan_kwh": float(record["base_natural"][t]), "final_contract_kwh": float(record["adjusted_natural"][t]), "adjustment_fee_yuan": float(record["penalty_natural"][t]), "charge_kwh": float(record["charge"][t]), "discharge_kwh": float(record["discharge"][t]), "soc_start_kwh": float(record["soc"][t]), "soc_end_kwh": float(record["soc"][t + 1]), "actual_net_kwh": float(inputs["net_actual_kwh"][index, t]), "emergency_kwh": float(record["emergency"][t]), "surplus_kwh": float(record["surplus"][t])})
    return rows


def locked_trajectory_audit(records: list[dict[str, Any]], releases: list[dict[str, Any]], inputs: dict[str, Any], start: int, updates: tuple[int, ...]) -> dict[str, Any]:
    """核实执行充放电量逐时段等于发布时锁定的共同储能轨迹，并统计紧急购电同时充电。"""
    released_issues = sorted({0} | {ISSUE_INTERVALS[i] for i in updates})
    index = {(entry["date"], entry["issue"]): entry for entry in releases}
    worst = 0.0
    overlap_intervals = 0
    overlap_energy = 0.0
    overlap_is_locked = True
    for record in records[start:]:
        iso = record["date"].isoformat()
        for t in range(144):
            governing = max(issue for issue in released_issues if issue <= t)
            entry = index[(iso, ISSUE_LABELS[ISSUE_INTERVALS.index(governing)])]
            worst = max(worst, abs(float(record["charge"][t]) - float(entry["common_charge_kwh"][t - governing])), abs(float(record["discharge"][t]) - float(entry["common_discharge_kwh"][t - governing])))
            if record["emergency"][t] > 1e-7 and record["charge"][t] > 1e-7:
                overlap_intervals += 1
                overlap_energy += float(record["emergency"][t])
                overlap_is_locked = overlap_is_locked and abs(float(record["charge"][t]) - float(entry["common_charge_kwh"][t - governing])) < 1e-4
    return {"max_abs_deviation_from_locked_trajectory_kwh": worst, "emergency_while_charging_intervals": overlap_intervals, "emergency_while_charging_energy_kwh": overlap_energy, "all_overlap_charges_equal_locked_trajectory": bool(overlap_is_locked)}


def assemble(primary: list[dict[str, Any]], policies: dict[str, list[dict[str, Any]]], releases: list[dict[str, Any]], inputs: dict[str, Any], cfg: Q3Config, all_releases: dict[str, list[dict[str, Any]]] | None = None) -> tuple[dict[str, Any], dict[str, Any], dict[str, list[dict[str, Any]]]]:
    start = inputs["dates"].index(DELIVERY_START)
    delivered = primary[start:]
    price, template_price = inputs["price_natural"], inputs["price_template"]
    comparison = {label: totals(records, start) for label, records in policies.items()}
    daily: list[dict[str, Any]] = []
    intervals = detail_rows(primary, inputs, start)
    days: list[dict[str, Any]] = []
    errors: dict[str, list[float]] = {label: [] for label in ISSUE_LABELS}
    max_balance = max_soc_residual = 0.0
    for index, record in enumerate(delivered, start=start):
        actual = inputs["net_actual_kwh"][index]
        balance = record["grid"] + record["emergency"] + record["discharge"] - record["charge"] - record["surplus"] - actual
        soc_residual = record["soc"][1:] - record["soc"][:-1] - cfg.eta_charge * record["charge"] + record["discharge"] / cfg.eta_discharge
        max_balance = max(max_balance, float(np.max(np.abs(balance))))
        max_soc_residual = max(max_soc_residual, float(np.max(np.abs(soc_residual))))
        for issue_i, forecast, observed in record["issue_metrics"]:
            errors[ISSUE_LABELS[issue_i]].extend(map(float, observed - forecast))
        charges, discharges = block_sums(record["charge"]), block_sums(record["discharge"])
        blocks = [{"time_range": f"{4 * i}:00-{4 * (i + 1)}:00", "charge_kwh": charges[i], "discharge_kwh": discharges[i]} for i in range(6)]
        template_penalty = record["penalty_template"]
        day = {
            "date": record["date"].isoformat(),
            "plan_kwh": list(map(float, record["plan"])),
            "adjusted_kwh": list(map(float, record["adjusted"])),
            "plan_total_kwh": float(record["plan"].sum()),
            "plan_cost_yuan": float(template_price @ record["plan"]),
            "adjusted_total_kwh": float(record["adjusted"].sum()),
            "adjustment_fee_template_yuan": float(template_penalty.sum()),
            "adjusted_settlement_yuan": float(template_price @ record["plan"] + template_penalty.sum()),
            "storage_blocks": blocks,
            "soc_start_kwh": float(record["soc"][0]),
            "soc_end_kwh": float(record["soc"][-1]),
            "emergency_segments": emergency_segments(record["emergency"]),
            "emergency_total_kwh": float(record["emergency"].sum()),
            "emergency_cost_yuan": float(record["emergency_cost"]),
            "natural_day_plan_kwh": float(record["base_natural"].sum()),
            "natural_day_final_contract_kwh": float(record["adjusted_natural"].sum()),
            "natural_day_settlement_yuan": float(record["settlement_cost"]),
            "residual_window_days": record["windows"],
            "scenario_counts": record["scenario_counts"],
        }
        days.append(day)
        daily.append({"policy": record["policy"], "date": day["date"], "plan_cost_yuan": day["plan_cost_yuan"], "adjustment_fee_yuan": float(record["penalty_natural"].sum()), "settlement_cost_yuan": record["settlement_cost"], "emergency_kwh": day["emergency_total_kwh"], "emergency_cost_yuan": record["emergency_cost"], "total_cost_yuan": record["settlement_cost"] + record["emergency_cost"], "soc_start_kwh": day["soc_start_kwh"], "soc_end_kwh": day["soc_end_kwh"]})
    fixed_refund = float(sum(price @ (record["adjusted_natural"] + 0.5 * np.abs(record["adjusted_natural"] - record["base_natural"])) + record["emergency_cost"] for record in delivered))
    # 主口径“原计划费保留＋逐笔调整费”下，下调已被支付的电量既要再付0.5p、又会抬高后续5p紧急购电风险，
    # 因此Q⁻=0在最优解中总成立。此时主口径费用恒等于“按最终净额退款”口径在同一固定电量上的结算值。
    # 这里把该恒等式的两个前提量显式记录，供报告与验收核对，避免把恒等误读为复制错误。
    structure = {label: {"increase_kwh": 0.0, "decrease_kwh": 0.0} for label in policies}
    if all_releases:
        for label, entries in all_releases.items():
            structure[label] = {
                "increase_kwh": float(sum(entry["increase_kwh"] for entry in entries if entry["date"] >= DELIVERY_START.isoformat())),
                "decrease_kwh": float(sum(entry["decrease_kwh"] for entry in entries if entry["date"] >= DELIVERY_START.isoformat())),
            }
    # 只有在逐次调整全部为上调（Q⁻≡0）时，主口径与“退款/最终净额”口径在固定电量上才严格相等。
    fixed_refund_equals_primary = bool(
        all(abs(item["decrease_kwh"]) < 1e-9 for label, item in structure.items() if label != "6+12+18_refund")
        and abs(fixed_refund - comparison["6+12+18"]["total_cost_yuan"]) < 1e-6
    )
    accuracy = {label: {"mae_kwh_per_interval": float(np.mean(np.abs(values))), "rmse_kwh_per_interval": float(np.sqrt(np.mean(np.asarray(values) ** 2))), "observations": len(values)} for label, values in errors.items()}
    charge = np.concatenate([record["charge"] for record in delivered])
    discharge = np.concatenate([record["discharge"] for record in delivered])
    emergency = np.concatenate([record["emergency"] for record in delivered])
    soc = np.concatenate([record["soc"] for record in delivered])
    audit = locked_trajectory_audit(primary, releases, inputs, start, (1, 2, 3))
    summary = {
        "model": "causal_rolling_SAA_MPC_with_common_storage_actions",
        "model_version": MODEL_VERSION,
        "strict_multistage_optimal": False,
        "period": {"start": "2025-02-01", "end": "2025-12-31", "days": 334},
        "time_convention": "区间起点；自然日00:00—24:00；模板行00:10—次日00:00；模板末列等于下一自然日00:00",
        "primary_settlement_assumption": "原计划费用保留；每次提交相对当前生效合同逐笔计费：上调1.5p、下调0.5p；取消后恢复会再次计费；需人工确认",
        "alternative_settlement_assumption": "退款且只按最终净额结算：pA+0.5p|A-G|；需人工确认；仅作敏感性，不与主口径混算",
        "totals_natural_day": comparison["6+12+18"],
        "policy_comparison": comparison,
        "fixed_quantity_refund_resettlement_total_cost_yuan": fixed_refund,
        "adjustment_structure": structure,
        "fixed_quantity_refund_resettlement_equals_primary": fixed_refund_equals_primary,
        "forecast_accuracy_next_block": accuracy,
        "release_audit": audit,
        "mapping_audit": inputs["mapping_audit"],
        "checks": {
            "max_energy_balance_residual_kwh": max_balance,
            "max_soc_residual_kwh": max_soc_residual,
            "soc_min_observed_kwh": float(soc.min()),
            "soc_max_observed_kwh": float(soc.max()),
            "max_interval_storage_energy_kwh": float(max(charge.max(), discharge.max())),
            "simultaneous_charge_discharge_intervals": int(np.count_nonzero((charge > 1e-7) & (discharge > 1e-7))),
            "emergency_while_charging_intervals": int(np.count_nonzero((emergency > 1e-7) & (charge > 1e-7))),
            "cross_day_soc_continuity": all(abs(delivered[i]["soc"][-1] - delivered[i + 1]["soc"][0]) < 1e-7 for i in range(len(delivered) - 1)),
        },
        "target_dates": {day["date"]: day for day in days if date.fromisoformat(day["date"]) in TARGET_DATES},
        "input_hashes_sha256": inputs["input_hashes"],
    }
    payload = {"metadata": {"model_version": MODEL_VERSION, "time_convention": summary["time_convention"], "settlement": summary["primary_settlement_assumption"], "input_hashes_sha256": inputs["input_hashes"]}, "days": days}
    return summary, payload, {"daily": daily, "interval": intervals}


def fmt(value: float) -> str:
    return f"{value:,.4f}"


def build_report(summary: dict[str, Any]) -> str:
    totals_block = summary["totals_natural_day"]
    comparison = summary["policy_comparison"]
    base = comparison["none"]["total_cost_yuan"]
    audit = summary["release_audit"]
    lines = [
        "# 第三问计算报告（因果滚动SAA/MPC）",
        "",
        "## 1. 模型定位与结算口径",
        "",
        "每天0:00制定当天计划，6:00、12:00、18:00仅调整尚未交付区间。同一发布时刻的全部情景共享购电量、充电量、放电量和SOC轨迹，只有紧急购电与弃电随情景变化，因此不存在情景提前预知未来；下一6小时严格执行求解器给出的共同储能轨迹，不再用事后贪心规则覆盖。该方法仍是滚动近似，`strict_multistage_optimal=false`，不宣称严格多阶段随机最优。",
        "",
        "主结算口径：初始计划费保留，每次提交相对当前生效合同逐笔计费，上调按1.5p、下调按0.5p；取消后恢复会产生两次费用。该合同解释标为**需人工确认**。退款且只按最终净额结算的口径仅作敏感性，不与主口径混算。",
        "",
        "| 指标（2025-02-01至12-31，主策略6+12+18） | 结果 |",
        "|---|---:|",
        f"| 计划购电量（自然日）/kWh | {fmt(totals_block['planned_purchase_kwh'])} |",
        f"| 最终合同购电量（自然日）/kWh | {fmt(totals_block['final_contract_kwh'])} |",
        f"| 非紧急结算费用/元 | {fmt(totals_block['settlement_cost_yuan'])} |",
        f"| 紧急购电量/kWh | {fmt(totals_block['emergency_kwh'])} |",
        f"| 紧急购电费用/元 | {fmt(totals_block['emergency_cost_yuan'])} |",
        f"| 总费用/元 | {fmt(totals_block['total_cost_yuan'])} |",
        "",
        "## 2. 8种预报组合比较",
        "",
        "每套策略独立维护SOC、跨日计划量和跨日最终合同量，均为全年连续回测。",
        "",
        "| 日内新预报 | 总费用/元 | 相对仅0:00节省/元 | 紧急购电量/kWh |",
        "|---|---:|---:|---:|",
    ]
    for label in ("none", "6", "12", "18", "6+12", "6+18", "12+18", "6+12+18"):
        item = comparison[label]
        lines.append(f"| {label} | {fmt(item['total_cost_yuan'])} | {fmt(base - item['total_cost_yuan'])} | {fmt(item['emergency_kwh'])} |")
    frozen, refund = comparison["6+12+18_frozen0"], comparison["6+12+18_refund"]
    structure = summary.get("adjustment_structure", {})
    main_structure = structure.get("6+12+18", {"increase_kwh": 0.0, "decrease_kwh": 0.0})
    refund_structure = structure.get("6+12+18_refund", {"increase_kwh": 0.0, "decrease_kwh": 0.0})
    lines += [
        "",
        f"冻结0:00光伏预报、但仍在6/12/18时重优化的基线：**{fmt(frozen['total_cost_yuan'])} 元**，用于区分新预报信息与单纯重优化机会。",
        "",
        "## 3. 结算敏感性",
        "",
        f"退款/最终净额口径重新优化：**{fmt(refund['total_cost_yuan'])} 元**；主策略固定电量仅改口径重算：**{fmt(summary['fixed_quantity_refund_resettlement_total_cost_yuan'])} 元**。两者均只作敏感性，不替代主口径。",
        "",
        f"主策略全年累计上调 {fmt(main_structure['increase_kwh'])} kWh、累计下调 {fmt(main_structure['decrease_kwh'])} kWh；"
        f"退款口径下分别为上调 {fmt(refund_structure['increase_kwh'])} kWh、下调 {fmt(refund_structure['decrease_kwh'])} kWh。",
        "",
        "主口径下 8 套策略的下调量全部为 0，这不是代码限制而是结算规则的经济后果：在“原计划费保留、不退款”下，"
        "减少已支付的电量既要再付 0.5p、又必然抬高后续按 5p 计价的紧急购电缺口，故恒被占优，最优解取 \\(Q^-\\equiv0\\)。"
        "此时对任一固定电量都有 \\(pA+0.5p|A-G| = pG+1.5pQ^+\\)，即“退款/最终净额”口径的结算值与原口径严格相等，"
        f"因此上表两个数字相同是恒等式的结果（校验标记 `fixed_quantity_refund_resettlement_equals_primary={summary.get('fixed_quantity_refund_resettlement_equals_primary')}`），"
        "而按退款口径**重新优化**得到的策略由于可以自由下调，费用明显低于主口径。",
        "",
        "## 4. 指定日期：表1计划购电量与最终调整购电量",
        "",
        "下表为交付口径的模板行，各行覆盖当天00:10至次日00:00；计划值来自该日0:00发布的计划，最终值来自当天最后一次调整后的生效合同。",
        "",
    ]
    for day in summary["target_dates"].values():
        lines += [f"### {day['date']}", "", "| 时间段 | 计划购电量/kWh | 最终调整购电量/kWh |", "|---|---:|---:|"]
        lines += [f"| {label} | {fmt(day['plan_kwh'][index])} | {fmt(day['adjusted_kwh'][index])} |" for label, index in WAIT_PERIODS]
        lines += [f"| 模板行计划合计 | {fmt(day['plan_total_kwh'])} | |", f"| 模板行最终合同合计 | | {fmt(day['adjusted_total_kwh'])} |", f"| 模板行计划购电费/元 | {fmt(day['plan_cost_yuan'])} | |", f"| 模板行调整费/元 | | {fmt(day['adjustment_fee_template_yuan'])} |", f"| 自然日结算费用/元 | | {fmt(day['natural_day_settlement_yuan'])} |", ""]
    lines += ["## 5. 指定日期：表2最终的4小时段充放电量", ""]
    for day in summary["target_dates"].values():
        lines += [f"### {day['date']}", "", "| 时间段 | 充电量/kWh | 放电量/kWh |", "|---|---:|---:|"]
        lines += [f"| {block['time_range']} | {fmt(block['charge_kwh'])} | {fmt(block['discharge_kwh'])} |" for block in day["storage_blocks"]]
        lines += [f"| 0:00储电量 | {fmt(day['soc_start_kwh'])} | |", f"| 24:00储电量 | {fmt(day['soc_end_kwh'])} | |", ""]
    lines += ["## 6. 指定日期：表3紧急购电", ""]
    for day in summary["target_dates"].values():
        lines += [f"### {day['date']}", ""]
        if day["emergency_segments"]:
            lines += ["| 时间段 | 购电量/kWh |", "|---|---:|"] + [f"| {segment['time_range']} | {fmt(segment['energy_kwh'])} |" for segment in day["emergency_segments"]]
        else:
            lines += ["无紧急购电。"]
        lines.append("")
    checks = summary["checks"]
    lines += [
        "## 7. 紧急购电同时充电的核查",
        "",
        f"主策略在{audit['emergency_while_charging_intervals']}个时段同时出现紧急购电与充电，涉及紧急购电量{fmt(audit['emergency_while_charging_energy_kwh'])} kWh。这些时段不是执行器自相矛盾：发布时锁定的共同储能轨迹在实际净负荷偏高时仍然被执行，缺口由紧急购电补足。",
        "",
        f"逐时段核对结果为：实际充电量与放电量相对发布日志中锁定共同轨迹的最大偏差为{audit['max_abs_deviation_from_locked_trajectory_kwh']:.3e} kWh，全部重叠时段的充电量均等于锁定轨迹（{audit['all_overlap_charges_equal_locked_trajectory']}）。未发现以贪心规则覆盖优化结果。",
        "",
        "## 8. 数值校验",
        "",
        "| 校验项 | 结果 |",
        "|---|---:|",
        f"| 最大能量平衡残差 | {checks['max_energy_balance_residual_kwh']:.3e} kWh |",
        f"| 最大SOC递推残差 | {checks['max_soc_residual_kwh']:.3e} kWh |",
        f"| SOC范围 | {fmt(checks['soc_min_observed_kwh'])}—{fmt(checks['soc_max_observed_kwh'])} kWh |",
        f"| 最大单时段充放电量 | {fmt(checks['max_interval_storage_energy_kwh'])} kWh |",
        f"| 同时充放电时段 | {checks['simultaneous_charge_discharge_intervals']} |",
        f"| 紧急购电同时充电时段 | {checks['emergency_while_charging_intervals']} |",
        f"| 跨日SOC连续 | {'是' if checks['cross_day_soc_continuity'] else '否'} |",
        "",
        f"时标核验：共核对{summary['mapping_audit']['actual_midnights_checked']}个自然日午夜，附件2跨行映射最大误差{summary['mapping_audit']['actual_midnight_max_error_kw']:.3e} kW；2025-02-01 00:00实际净负荷为{fmt(summary['mapping_audit']['feb1_midnight_net_kwh'])} kWh。",
        "",
        "## 9. 尚存限制",
        "",
        "1. 0:00计划没有用完整情景树联合定价未来6/12/18时的信息到达与调整机会，属于滚动SAA/MPC近似。",
        "2. 终端储能价值固定为0.9×最低电价，未表达次日清晨负荷与紧急电价风险，该系数不是由最优性推导唯一确定。",
        "3. 附件3只给整点光伏预报，10分钟值是线性插值；发布时刻的插值首节点使用上一区间实际值，属于持续性近似。",
        "4. 附件未提供负荷预报，日内负荷预测沿用计划日前的历史轮廓，未用当天已观测偏差校正剩余时段。",
        "5. 午夜购电锁定为前一日最后一次提交的次日00:00量，属于合同时间假设。",
        "6. 结算口径（逐次提交不退款、退款/最终净额）均需人工确认。",
        "",
    ]
    return "\n".join(lines)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--data-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--report", type=Path, required=True)
    parser.add_argument("--checkpoint-dir", type=Path, default=None)
    args = parser.parse_args()
    cfg = Q3Config()
    output_dir = args.output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    args.report.parent.mkdir(parents=True, exist_ok=True)
    checkpoint_dir = (args.checkpoint_dir or (output_dir / "checkpoints")).resolve()
    checkpoint_dir.mkdir(parents=True, exist_ok=True)
    inputs = load_q3_inputs(args.data_dir.resolve(), cfg)
    signature = build_signature(cfg, inputs)
    cache = ForecastCache(inputs, cfg)
    policies: dict[str, list[dict[str, Any]]] = {}
    releases_by_policy: dict[str, list[dict[str, Any]]] = {}
    labels = [label for label, _, _, _ in policy_specs()]
    for label, updates, forecast_mode, settlement in policy_specs():
        path = checkpoint_dir / f"{label}.json"
        data = load_checkpoint(path, signature)
        if data is None:
            print(f"[计算] {label}", flush=True)
            records, releases = simulate_policy(inputs, cfg, updates, forecast_mode, settlement, cache)
            save_checkpoint(path, signature, label, records, releases)
            print(f"[完成] {label}", flush=True)
        else:
            print(f"[复用检查点] {label}", flush=True)
            records, releases = deserialize_records(data["days"]), data["releases"]
        policies[label] = records
        releases_by_policy[label] = releases
    releases = releases_by_policy["6+12+18"]
    summary, payload, tables = assemble(policies["6+12+18"], policies, releases, inputs, cfg, releases_by_policy)
    start = inputs["dates"].index(DELIVERY_START)
    (output_dir / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    (output_dir / "solver_payload.json").write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    (output_dir / "release_log.json").write_text(json.dumps([entry for entry in releases if entry["date"] >= DELIVERY_START.isoformat()], ensure_ascii=False), encoding="utf-8")
    release_dir = output_dir / "release_logs"
    release_dir.mkdir(exist_ok=True)
    for label in labels:
        (release_dir / f"{label}.json").write_text(json.dumps([entry for entry in releases_by_policy[label] if entry["date"] >= DELIVERY_START.isoformat()], ensure_ascii=False), encoding="utf-8")
    write_csv(output_dir / "daily_metrics.csv", tables["daily"])
    write_csv(output_dir / "interval_detail.csv", tables["interval"])
    comparison = []
    for label in labels:
        for record in policies[label][start:]:
            comparison.append({"policy": label, "date": record["date"].isoformat(), "settlement_cost_yuan": record["settlement_cost"], "emergency_kwh": float(record["emergency"].sum()), "emergency_cost_yuan": record["emergency_cost"], "total_cost_yuan": record["settlement_cost"] + record["emergency_cost"], "soc_start_kwh": float(record["soc"][0]), "soc_end_kwh": float(record["soc"][-1])})
    write_csv(output_dir / "policy_comparison_daily.csv", comparison)
    write_csv(output_dir / "baseline_none_interval_detail.csv", detail_rows(policies["none"], inputs, start))
    write_csv(output_dir / "baseline_frozen0_interval_detail.csv", detail_rows(policies["6+12+18_frozen0"], inputs, start))
    args.report.write_text(build_report(summary), encoding="utf-8")
    print(json.dumps(summary["totals_natural_day"], ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
