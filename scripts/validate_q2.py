from __future__ import annotations

import csv
import json
from pathlib import Path

import numpy as np
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "outputs" / "q2"


def main() -> None:
    summary = json.loads((OUTPUT / "summary.json").read_text(encoding="utf-8"))
    payload = json.loads((OUTPUT / "solver_payload.json").read_text(encoding="utf-8"))
    assert len(payload["days"]) == 334
    assert summary["period"] == {"start": "2025-02-01", "end": "2025-12-31", "days": 334}
    assert abs(summary["mapping_audit"]["feb1_midnight_net_kwh"] - 423.96) < 1e-9
    assert summary["mapping_audit"]["max_abs_error_kwh"] < 1e-10

    workbook = load_workbook(OUTPUT / "result2.xlsx", read_only=False, data_only=True)
    plan = workbook["计划购电量"]
    assert plan.cell(335, 1).value is not None and plan.cell(335, 147).value is not None
    assert plan.cell(336, 1).value is None
    plan_rows = list(plan.iter_rows(min_row=2, max_row=335, min_col=2, max_col=147, values_only=True))
    for workbook_row, day in zip(plan_rows, payload["days"]):
        workbook_values = np.asarray(workbook_row, dtype=float)
        expected = np.asarray(day["template_grid_kwh"] + [day["template_grid_total_kwh"], day["template_grid_cost_yuan"]])
        assert np.allclose(workbook_values, expected, atol=1e-8, rtol=0)

    storage = workbook["充放电量"]
    workbook_charge = sum(float(storage.cell(row, 3).value or 0) for row in range(2, 2006))
    workbook_discharge = sum(float(storage.cell(row, 4).value or 0) for row in range(2, 2006))
    expected_charge = sum(block["charge_kwh"] for day in payload["days"] for block in day["storage_blocks"])
    expected_discharge = sum(block["discharge_kwh"] for day in payload["days"] for block in day["storage_blocks"])
    assert abs(workbook_charge - expected_charge) < 1e-7
    assert abs(workbook_discharge - expected_discharge) < 1e-7

    emergency = workbook["紧急购电量"]
    emergency_rows = 1 + sum(max(1, len(day["emergency_segments"])) for day in payload["days"])
    workbook_emergency = sum(float(emergency.cell(row, 3).value or 0) for row in range(2, emergency_rows + 1))
    assert emergency.cell(emergency_rows, 3).value is not None
    assert abs(workbook_emergency - summary["totals_natural_day"]["emergency_purchase_kwh"]) < 1e-7

    with (OUTPUT / "interval_detail.csv").open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    assert len(rows) == 334 * 144
    assert not any(float(row["emergency_kwh"]) > 1e-7 and float(row["charge_kwh"]) > 1e-7 for row in rows)
    checks = summary["checks"]
    assert checks["max_energy_balance_residual_kwh"] < 1e-8
    assert checks["max_soc_residual_kwh"] < 1e-8
    assert checks["simultaneous_charge_discharge_intervals"] == 0
    assert checks["emergency_while_charging_intervals"] == 0
    assert checks["cross_day_soc_continuity"] is True
    print("Q2 validation passed: mapping, 48,096 intervals, constraints, and workbook reconciliation.")


if __name__ == "__main__":
    main()
