"""第三问独立验收：逐项核对保存结果、发布日志、结算口径、物理约束与工作簿。"""
from __future__ import annotations

import csv
import json
import re
from functools import lru_cache
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs" / "q3"
DELIVERY_START = "2025-02-01"
ISSUE_LABELS = ("0:00", "6:00", "12:00", "18:00")
ISSUE_INTERVALS = (0, 36, 72, 108)
LABELS = ("none", "6", "12", "18", "6+12", "6+18", "12+18", "6+12+18", "6+12+18_frozen0", "6+12+18_refund")
DELIVERED_DAYS = 334
ETA = 0.9
# 储能功率上限 5000 kW，区间长度 1/6 h，故每时段能量上限为 833.3333…kWh。
# 四位小数写法 833.3333 是显示舍入值，不能当作精确上界，判据必须用精确限值。
INTERVAL_LIMIT_KWH = 5000.0 / 6.0


def read_csv(name: str) -> list[dict[str, str]]:
    with (OUT / name).open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def updates_of(label: str) -> tuple[int, ...]:
    base = label.replace("_frozen0", "").replace("_refund", "")
    if base == "none":
        return ()
    return tuple(ISSUE_INTERVALS[ISSUE_LABELS.index(f"{int(hour)}:00")] for hour in base.split("+"))


def check(condition: bool, message: str) -> None:
    if not condition:
        raise AssertionError(message)


@lru_cache(maxsize=None)
def checkpoint(label: str) -> dict:
    return json.loads((OUT / "checkpoints" / f"{label}.json").read_text(encoding="utf-8"))


@lru_cache(maxsize=None)
def release_entries(label: str) -> dict[tuple[str, str], dict]:
    entries = json.loads((OUT / "release_logs" / f"{label}.json").read_text(encoding="utf-8"))
    return {(entry["date"], entry["issue"]): entry for entry in entries}


def delivered_days(label: str) -> list[dict]:
    return checkpoint(label)["days"][-DELIVERED_DAYS:]


def main() -> None:
    summary = json.loads((OUT / "summary.json").read_text(encoding="utf-8"))
    payload = json.loads((OUT / "solver_payload.json").read_text(encoding="utf-8"))
    releases = json.loads((OUT / "release_log.json").read_text(encoding="utf-8"))
    check(summary["strict_multistage_optimal"] is False, "必须声明不是严格多阶段随机最优")
    check(summary["model"] == "causal_rolling_SAA_MPC_with_common_storage_actions", "模型名称应为因果滚动SAA/MPC")

    # 1. 时标与跨行映射
    mapping = summary["mapping_audit"]
    check(abs(mapping["feb1_midnight_net_kwh"] - 423.96) < 1e-9, f"2025-02-01 00:00净负荷应为423.9600，实为{mapping['feb1_midnight_net_kwh']}")
    check(mapping["actual_midnight_max_error_kw"] < 1e-10, "自然日午夜跨行映射存在误差")
    check("00:10" in mapping["template_row"] and "00:00" in mapping["natural_day"], "时间口径说明缺失")

    # 2. 交付期规模与10套策略
    check(len(payload["days"]) == DELIVERED_DAYS, "载荷应有334天")
    check(set(summary["policy_comparison"]) == set(LABELS), "策略集合不完整")
    comparison = read_csv("policy_comparison_daily.csv")
    check(len(comparison) == 334 * 10, f"策略日记录应为3,340条，实为{len(comparison)}")
    for label in LABELS:
        rows = [row for row in comparison if row["policy"] == label]
        check(len(rows) == DELIVERED_DAYS, f"{label}应有334天，实为{len(rows)}")
        check(all(abs(float(row["total_cost_yuan"]) - float(row["settlement_cost_yuan"]) - float(row["emergency_cost_yuan"])) < 1e-6 for row in rows), f"{label}总费用不等于结算费加紧急购电费")
        check(all(1200 - 1e-6 <= float(row["soc_start_kwh"]) <= 10800 + 1e-6 and 1200 - 1e-6 <= float(row["soc_end_kwh"]) <= 10800 + 1e-6 for row in rows), f"{label}的SOC越界")
        check(all(abs(float(rows[i - 1]["soc_end_kwh"]) - float(rows[i]["soc_start_kwh"])) < 1e-7 for i in range(1, DELIVERED_DAYS)), f"{label}跨日SOC不连续")

    # 3. 主策略明细与费用独立复算
    rows = read_csv("interval_detail.csv")
    check(len(rows) == DELIVERED_DAYS * 144, f"主策略明细应为48,096条，实为{len(rows)}")
    price = np.array([float(row["price_yuan_per_kwh"]) for row in rows])
    plan = np.array([float(row["initial_plan_kwh"]) for row in rows])
    final = np.array([float(row["final_contract_kwh"]) for row in rows])
    fee = np.array([float(row["adjustment_fee_yuan"]) for row in rows])
    emergency = np.array([float(row["emergency_kwh"]) for row in rows])
    charge = np.array([float(row["charge_kwh"]) for row in rows])
    discharge = np.array([float(row["discharge_kwh"]) for row in rows])
    soc_start = np.array([float(row["soc_start_kwh"]) for row in rows])
    soc_end = np.array([float(row["soc_end_kwh"]) for row in rows])
    actual = np.array([float(row["actual_net_kwh"]) for row in rows])
    surplus = np.array([float(row["surplus_kwh"]) for row in rows])
    totals = summary["totals_natural_day"]
    settlement = float(price @ plan + fee.sum())
    emergency_cost = float(5 * price @ emergency)
    check(abs(settlement - totals["settlement_cost_yuan"]) < 1e-5, f"非紧急结算费不可复算：{settlement} vs {totals['settlement_cost_yuan']}")
    check(abs(emergency_cost - totals["emergency_cost_yuan"]) < 1e-5, f"紧急购电费不可复算：{emergency_cost} vs {totals['emergency_cost_yuan']}")
    check(abs(totals["total_cost_yuan"] - settlement - emergency_cost) < 1e-5, "主口径总费用不等于计划费加逐次调整费加紧急购电费")
    check(abs(totals["emergency_kwh"] - emergency.sum()) < 1e-6, "紧急购电量不一致")

    # 4. 物理约束与能量平衡
    balance = final + discharge + emergency - charge - surplus - actual
    soc_residual = soc_end - soc_start - ETA * charge + discharge / ETA
    check(float(np.max(np.abs(balance))) < 1e-7, f"能量平衡最大残差{np.max(np.abs(balance)):.3e} kWh超限")
    check(float(np.max(np.abs(soc_residual))) < 1e-7, f"SOC递推最大残差{np.max(np.abs(soc_residual)):.3e} kWh超限")
    check(soc_start.min() >= 1200 - 1e-9 and soc_end.max() <= 10800 + 1e-9, "SOC越界")
    peak_storage = max(charge.max(), discharge.max())
    check(peak_storage <= INTERVAL_LIMIT_KWH + 1e-6, f"单时段充放电超过限值{INTERVAL_LIMIT_KWH:.7f} kWh，实测峰值{peak_storage:.7f} kWh")
    check(int(np.count_nonzero((charge > 1e-7) & (discharge > 1e-7))) == 0, "存在同时充放电")
    checks = summary["checks"]
    check(checks["max_energy_balance_residual_kwh"] < 1e-7 and checks["max_soc_residual_kwh"] < 1e-7, "汇总残差超限")
    check(checks["cross_day_soc_continuity"] is True, "跨日SOC不连续")

    # 5. 发布日志：完整合同、共同储能向量、信息截止时刻、调整费用
    check(len(releases) == DELIVERED_DAYS * 4, f"主策略发布日志应有1,336条，实为{len(releases)}")
    for entry in releases:
        expected = len(entry["contract_kwh"]) + (1 if entry["issue"] == "0:00" else 0)
        check(len(entry["common_charge_kwh"]) == expected == len(entry["common_discharge_kwh"]), "发布日志向量长度不一致")
        check(bool(entry["information_cutoff"]) and entry["solve_status"] == "optimal" and entry["scenario_count"] >= 1, "发布日志缺少截止时刻、求解状态或情景数")
        check(np.isfinite(entry["adjustment_fee_yuan"]) and entry["adjustment_fee_yuan"] >= 0, "调整费用不合法")
    check(all(entry["increase_kwh"] == 0.0 and entry["decrease_kwh"] == 0.0 for entry in releases if entry["issue"] == "0:00"), "0:00发布不应产生调整量")
    check(sum(entry["adjustment_fee_yuan"] for entry in releases) > 0, "调整费用不应恒为零")

    # 6. 执行量必须等于发布时锁定的共同储能轨迹（不得用贪心补救覆盖优化结果）
    worst_locked = 0.0
    overlap = 0
    overlap_energy = 0.0
    overlap_locked = True
    for label in LABELS:
        entries = release_entries(label)
        released = sorted({0} | set(updates_of(label)))
        locked_charge = {issue: 0.0 for issue in released}
        for day in delivered_days(label):
            iso = day["date"]
            for slot in range(144):
                governing = max(issue for issue in released if issue <= slot)
                entry = entries[(iso, ISSUE_LABELS[ISSUE_INTERVALS.index(governing)])]
                deviation = max(abs(day["charge"][slot] - entry["common_charge_kwh"][slot - governing]), abs(day["discharge"][slot] - entry["common_discharge_kwh"][slot - governing]))
                worst_locked = max(worst_locked, deviation)
                if day["emergency"][slot] > 1e-7 and day["charge"][slot] > 1e-7:
                    overlap_locked = overlap_locked and deviation < 1e-4
                if label == "6+12+18" and day["emergency"][slot] > 1e-7 and day["charge"][slot] > 1e-7:
                    overlap += 1
                    overlap_energy += float(day["emergency"][slot])
    check(worst_locked < 1e-4, f"执行充放电与发布锁定轨迹最大偏差{worst_locked:.3e} kWh")
    audit = summary["release_audit"]
    check(audit["emergency_while_charging_intervals"] == overlap, "汇总与逐日明细的紧急购电同时充电时段数不一致")
    check(abs(audit["emergency_while_charging_energy_kwh"] - overlap_energy) < 1e-6, "紧急购电同时充电电量不一致")
    check(audit["all_overlap_charges_equal_locked_trajectory"] is True and overlap_locked, "存在紧急购电同时充电时充电量不等于锁定轨迹")

    # 7. 逐套策略独立复算结算费；退款口径不得与主口径混算
    day_price = price[:144]
    for label in LABELS:
        for day in delivered_days(label):
            base = np.asarray(day["base_natural"], dtype=float)
            final_vector = np.asarray(day["adjusted_natural"], dtype=float)
            fee_vector = np.asarray(day["penalty_natural"], dtype=float)
            if label.endswith("_refund"):
                check(float(np.abs(fee_vector).max()) < 1e-12, "退款口径不得累加逐次调整费")
                expected = float(day_price @ (final_vector + 0.5 * np.abs(final_vector - base)))
            else:
                expected = float(day_price @ base + fee_vector.sum())
            check(abs(expected - day["settlement_cost"]) < 1e-6, f"{label} {day['date']} 结算费不可复算")
        if label == "none":
            check(all(abs(value) < 1e-12 for day in delivered_days(label) for value in day["penalty_template"]), "仅0:00策略不应产生调整费")
    check(abs(totals["total_cost_yuan"] - summary["policy_comparison"]["6+12+18"]["total_cost_yuan"]) < 1e-9, "主口径汇总与主策略不一致")
    check(abs(totals["total_cost_yuan"] - summary["policy_comparison"]["6+12+18_refund"]["total_cost_yuan"]) > 1.0, "退款口径被混入主口径")

    # 7b. 主口径下 Q-≡0 的推论：固定电量按退款口径重算必须等于主口径总费用，且该恒等式可由明细独立复算。
    refund_on_fixed = float(np.sum(price * (final + 0.5 * np.abs(final - plan))) + emergency_cost)
    check(abs(refund_on_fixed - totals["total_cost_yuan"]) < 1e-5, f"固定电量退款口径重算与主口径不符：{refund_on_fixed} vs {totals['total_cost_yuan']}")
    check(abs(refund_on_fixed - summary["fixed_quantity_refund_resettlement_total_cost_yuan"]) < 1e-5, "汇总中的固定电量退款重算值不可由明细复算")
    check(summary["fixed_quantity_refund_resettlement_equals_primary"] is True, "固定电量退款重算等于主口径的标记应为真")
    structure = summary["adjustment_structure"]
    for label in LABELS:
        if not label.endswith("_refund"):
            check(abs(structure[label]["decrease_kwh"]) < 1e-9, f"{label} 在主口径下不应出现下调")
    check(structure["6+12+18_refund"]["decrease_kwh"] > 1.0, "退款口径应能自由下调，否则 Q-≡0 无法排除为代码限制")
    logged_decrease = sum(entry["decrease_kwh"] for entry in release_entries("6+12+18_refund").values())
    check(abs(structure["6+12+18_refund"]["decrease_kwh"] - logged_decrease) < 1e-6, "退款口径下调量与发布日志不符")

    # 8. 基线明细可独立复算
    for name, label in (("baseline_none_interval_detail.csv", "none"), ("baseline_frozen0_interval_detail.csv", "6+12+18_frozen0")):
        baseline = read_csv(name)
        check(len(baseline) == DELIVERED_DAYS * 144, f"{name}条数不足")
        recomputed = sum(5 * float(row["price_yuan_per_kwh"]) * float(row["emergency_kwh"]) + float(row["price_yuan_per_kwh"]) * float(row["initial_plan_kwh"]) + float(row["adjustment_fee_yuan"]) for row in baseline)
        check(abs(recomputed - summary["policy_comparison"][label]["total_cost_yuan"]) < 1e-4, f"{label}基线费用不可独立复算")

    # 9. 工作簿与载荷逐项一致，且无错误值
    workbook = load_workbook(OUT / "result3.xlsx", data_only=True)
    check(set(workbook.sheetnames) == {"计划购电量", "调整购电量", "充放电量", "紧急购电量"}, f"工作表名称异常：{workbook.sheetnames}")
    for sheet, key, total_key, cost_key in (("计划购电量", "plan_kwh", "plan_total_kwh", "plan_cost_yuan"), ("调整购电量", "adjusted_kwh", "adjusted_total_kwh", "adjusted_settlement_yuan")):
        values = list(workbook[sheet].iter_rows(min_row=2, max_row=335, min_col=2, max_col=147, values_only=True))
        check(len(values) == DELIVERED_DAYS, f"{sheet}行数不为334")
        for row, day in zip(values, payload["days"]):
            check(np.allclose(np.asarray(row, dtype=float), np.asarray(day[key] + [day[total_key], day[cost_key]]), atol=1e-8, rtol=0), f"{sheet} {day['date']} 与载荷不一致")
    storage = workbook["充放电量"]
    check(abs(sum(float(storage.cell(i, 3).value or 0) for i in range(2, 2006)) - sum(block["charge_kwh"] for day in payload["days"] for block in day["storage_blocks"])) < 1e-6, "充放电量表充电量不一致")
    check(abs(sum(float(storage.cell(i, 4).value or 0) for i in range(2, 2006)) - sum(block["discharge_kwh"] for day in payload["days"] for block in day["storage_blocks"])) < 1e-6, "充放电量表放电量不一致")
    for position, day in enumerate(payload["days"]):
        # 第1行为表头，故第 position 天占据 2+6*position 至 7+6*position 行；0:00 在首行块，24:00 在次行块。
        check(abs(float(storage.cell(2 + 6 * position, 6).value or 0) - day["soc_start_kwh"]) < 1e-8, f"充放电量表{day['date']} 0:00储电量不一致")
        check(abs(float(storage.cell(3 + 6 * position, 6).value or 0) - day["soc_end_kwh"]) < 1e-8, f"充放电量表{day['date']} 24:00储电量不一致")
    emergency_sheet = workbook["紧急购电量"]
    emergency_rows = 1 + sum(max(1, len(day["emergency_segments"])) for day in payload["days"])
    check(abs(sum(float(emergency_sheet.cell(i, 3).value or 0) for i in range(2, emergency_rows + 1)) - totals["emergency_kwh"]) < 1e-6, "紧急购电量表总量不一致")
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            for value in row:
                check(not (isinstance(value, str) and re.match(r"^#(REF|DIV/0|VALUE|NAME|N/A|NUM|NULL|SPILL|CALC)", value)), f"{sheet.title}存在错误值{value}")

    # 10. 交付期口径
    check(all(row["date"] >= DELIVERY_START for row in rows), "主策略明细早于交付期")
    check(len({row["date"] for row in rows}) == DELIVERED_DAYS, "主策略明细日期数不为334")
    check(len({row["date"] for row in comparison}) == DELIVERED_DAYS, "策略比较日期数不为334")

    print("Q3验收通过：10套策略各334天，共3,340条策略日记录；主策略48,096条自然日明细。")
    print(f"  2025-02-01 00:00实际净负荷 {mapping['feb1_midnight_net_kwh']:.4f} kWh；原始附件跨行映射最大误差 {mapping['actual_midnight_max_error_kw']:.3e} kW")
    print(f"  能量平衡最大残差 {checks['max_energy_balance_residual_kwh']:.3e} kWh；SOC递推最大残差 {checks['max_soc_residual_kwh']:.3e} kWh")
    print(f"  SOC范围 {checks['soc_min_observed_kwh']:.6f}—{checks['soc_max_observed_kwh']:.6f} kWh（数值容差内不越界）")
    print(f"  最大单时段充放电 {peak_storage:.7f} kWh，限值 {INTERVAL_LIMIT_KWH:.7f} kWh，裕度 {INTERVAL_LIMIT_KWH - peak_storage:+.3e} kWh")
    print(f"  同时充放电 {checks['simultaneous_charge_discharge_intervals']} 个时段；紧急购电同时充电 {checks['emergency_while_charging_intervals']} 个时段（{audit['emergency_while_charging_energy_kwh']:.4f} kWh）")
    print(f"  主策略总费用 {totals['total_cost_yuan']:.4f} 元（结算 {settlement:.4f} 元 + 紧急购电 {emergency_cost:.4f} 元）")
    print(f"  执行充放电与发布锁定共同轨迹最大偏差 {worst_locked:.3e} kWh；工作簿四张表与载荷逐项一致，无错误值")


if __name__ == "__main__":
    main()
