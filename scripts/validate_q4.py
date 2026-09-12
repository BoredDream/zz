"""第四问（波动电价两阶段随机规划）的独立验收。

两块内容：
  A. 物理与费用口径：逐时段核对 SOC 递推、功率平衡、充放电上限、SOC 上下限、非负性，
     验证"计划+调整+紧急"三项费用之和恒等于模型总费用；并核对结算用的是【实际】电价。
  B. 工作簿对账：把 result4-2.xlsx / result4-3.xlsx 的每一格与求解器落盘的 payload 逐项比对。

用法： python scripts/validate_q4.py <2|3> [K]
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import numpy as np
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))
import q4_solver as M  # noqa: E402
import q3_multistage as Q3  # noqa: E402

OUT = ROOT / "outputs" / "q4"
TOL = 1e-6
failures: list[str] = []


def check(name: str, worst: float, tol: float = TOL) -> None:
    ok = worst <= tol
    print(f"  [{'OK ' if ok else 'FAIL'}] {name:<42} 最大偏差 {worst:.3e}")
    if not ok:
        failures.append(f"{name}: {worst:.3e}")


def main() -> int:
    which = sys.argv[1] if len(sys.argv) > 1 else "3"
    K = sys.argv[2] if len(sys.argv) > 2 else "30"
    assert which in ("2", "3"), "variant 必须是 2 或 3"
    payload = json.loads((OUT / f"payload_q4-{which}_K{K}.json").read_text(encoding="utf-8"))
    detail = np.load(OUT / f"detail_q4-{which}_K{K}.npz")
    dates = [str(d) for d in detail["dates"]]
    x, q, z, c, g, w, S = (detail[k] for k in ("x", "q", "z", "c", "g", "w", "S"))
    S0 = detail["S0"]
    days = payload["days"]
    first = dates.index(days[0]["date"])

    print("=" * 88)
    print(f"A. 物理与费用口径  变体 4-{which}，K={K}，{len(dates)} 天 × {M.T} 时段")
    print("=" * 88)
    assert len(days) == 334 and days[0]["date"] == "2025-02-01" and dates[-1] == "2025-12-31", \
        f"交付期应为 2025-02-01..12-31，实得 {days[0]['date']}..{dates[-1]}"
    assert len(dates) == M.ND, f"全年应为 {M.ND} 天，实得 {len(dates)}"

    worst = {k: 0.0 for k in ("soc", "balance", "charge_limit", "discharge_limit",
                              "soc_min", "soc_max", "negativity", "cost_split", "soc_start_chain")}
    for i, d in enumerate(dates):
        prev = float(S0[i])
        for t in range(M.T):
            now = float(S[i, t])
            worst["soc"] = max(worst["soc"], abs(now - (prev + M.ETA * c[i, t] - g[i, t] / M.ETA)))
            worst["balance"] = max(worst["balance"],
                                   abs(q[i, t] + z[i, t] + g[i, t] - c[i, t] - w[i, t]
                                       - (M.L[i, t] - M.G[i, t])))
            worst["charge_limit"] = max(worst["charge_limit"], c[i, t] - M.CMAX)
            worst["discharge_limit"] = max(worst["discharge_limit"], g[i, t] - M.CMAX)
            worst["soc_min"] = max(worst["soc_min"], M.SMIN - now)
            worst["soc_max"] = max(worst["soc_max"], now - M.SMAX)
            worst["negativity"] = max(worst["negativity"], -min(c[i, t], g[i, t], z[i, t], w[i, t]))
            prev = now
        if i > 0:
            worst["soc_start_chain"] = max(worst["soc_start_chain"], abs(float(S0[i]) - float(S[i - 1, M.T - 1])))
        total = M.day_cost4(i, x[i], q[i], {"z": z[i]})[0]
        p1, p2, p3 = M.settle_parts4(i, x[i], q[i], z[i])
        worst["cost_split"] = max(worst["cost_split"], abs(total - (p1.sum() + p2.sum() + p3.sum())))
    for name, value in worst.items():
        check(name, float(value))
    check("充放电量上限 CMAX = 5000*DT", abs(M.CMAX - 5000.0 / 6.0), 1e-12)

    # 结算价必须是当日【实际】电价，而不是任何预测值
    bad_price = max(float(np.abs(detail["price"][i] - M.PMAT[i]).max()) for i in range(len(dates)))
    check("落盘电价 = 附件4 实际电价", bad_price)

    # 变体 4-2 不含可调整购电量，q 必须恒等于 x
    gap = float(np.abs(q - x).max())
    if which == "2":
        check("变体4-2 不可调整：max|q-x|", gap)
    else:
        print(f"  [info] 变体4-3 可调整：Σ|q-x| = {np.abs(q[first:]-x[first:]).sum():,.1f} kWh（交付期）")

    print(f"  [info] 含紧急购电的日期 {sum(1 for day in days if day['emergency_segments'])}/334")

    print()
    print("=" * 88)
    print(f"B. 工作簿 result4-{which}.xlsx 逐格对账")
    print("=" * 88)
    book = OUT / f"result4-{which}.xlsx"
    if not book.exists():
        print(f"  [FAIL] 未找到 {book}")
        failures.append("missing workbook")
        return 1
    wb = openpyxl.load_workbook(book)
    need = ["计划购电量"] + (["调整购电量"] if which == "3" else []) + ["充放电量", "紧急购电量"]
    assert wb.sheetnames == need, f"表结构应为 {need}，实得 {wb.sheetnames}"

    def price_sheet(name: str, kwh: str, total: str, cost: str) -> None:
        ws = wb[name]
        bad_kwh = bad_total = bad_cost = bad_date = 0.0
        for r, day in enumerate(days, start=2):
            if ws.cell(r, 1).value.date().isoformat() != day["date"]:
                bad_date += 1
            for col, value in enumerate(day[kwh], start=2):
                bad_kwh = max(bad_kwh, abs(float(ws.cell(r, col).value or 0.0) - value))
            bad_total = max(bad_total, abs(float(ws.cell(r, 146).value or 0.0) - day[total]))
            bad_cost = max(bad_cost, abs(float(ws.cell(r, 147).value or 0.0) - day[cost]))
        check(f"{name} 日期列", float(bad_date))
        check(f"{name} 144列购电量", bad_kwh)
        check(f"{name} 全天购电量", bad_total)
        check(f"{name} 全天购电费", bad_cost)

    price_sheet("计划购电量", "plan_kwh", "plan_total_kwh", "plan_cost_yuan")
    if which == "3":
        price_sheet("调整购电量", "adjusted_kwh", "adjusted_total_kwh", "adjusted_cost_yuan")

    ws = wb["充放电量"]
    bad = {k: 0.0 for k in ("rows", "block", "charge", "discharge", "time", "soc")}
    expected_labels = [b["time_range"] for b in days[0]["storage_blocks"]]
    for i, day in enumerate(days):
        if ws.cell(2 + 6 * i, 1).value is None:
            bad["rows"] += 1
        for k, block in enumerate(day["storage_blocks"]):
            r = 2 + 6 * i + k
            bad["block"] = max(bad["block"], 0.0 if ws.cell(r, 2).value == block["time_range"] else 1.0)
            bad["charge"] = max(bad["charge"], abs(float(ws.cell(r, 3).value or 0.0) - block["charge_kwh"]))
            bad["discharge"] = max(bad["discharge"], abs(float(ws.cell(r, 4).value or 0.0) - block["discharge_kwh"]))
        bad["soc"] = max(bad["soc"],
                         abs(float(ws.cell(2 + 6 * i, 6).value or 0.0) - day["soc_start_kwh"]),
                         abs(float(ws.cell(3 + 6 * i, 6).value or 0.0) - day["soc_end_kwh"]))
        bad["time"] = max(bad["time"],
                          0.0 if ws.cell(2 + 6 * i, 5).value == "0:10" else 1.0,
                          0.0 if ws.cell(3 + 6 * i, 5).value == "0:10+1" else 1.0)
    print(f"  [info] 表2 段标签 {expected_labels}")
    check("表2 日期行结构(334×6行)", float(bad["rows"]))
    check("表2 六个段标签", bad["block"])
    check("表2 充电量", bad["charge"])
    check("表2 放电量", bad["discharge"])
    check("表2 时刻列", bad["time"])
    check("表2 0:10 / 次日0:10 储电量", bad["soc"])

    ws = wb["紧急购电量"]
    expect = [row for day in days for row in
              ([(day["date"], seg["time_range"], seg["energy_kwh"]) for seg in day["emergency_segments"]]
               or [(day["date"], None, 0.0)])]
    bad_e = bad_e_span = 0.0
    for r, (dt, span, kwh) in enumerate(expect, start=2):
        got_span = ws.cell(r, 2).value
        if (got_span is None) != (span is None) or (span is not None and got_span != span):
            bad_e_span += 1
        bad_e = max(bad_e, abs(float(ws.cell(r, 3).value or 0.0) - kwh))
    check("表2 紧急购电时段标签", float(bad_e_span))
    check("表2 紧急购电量", bad_e)
    check("表2 紧急购电行数", float(abs(ws.max_row - (len(expect) + 1))))
    seg_sum = sum(seg["energy_kwh"] for day in days for seg in day["emergency_segments"])
    day_sum = sum(day["emergency_total_kwh"] for day in days)
    check("表2 紧急购电：分段之和 vs 逐日合计", abs(seg_sum - day_sum), 1e-6)
    check("表2 紧急购电：逐日合计 vs 汇总", abs(day_sum - payload["meta"]["totals"]["emergency_kwh"]), 1e-6)
    check("表2 紧急购电：工作簿 vs 求解器原始数组", abs(day_sum - float(z[first:].sum())), 1e-6)

    # 汇总口径：三项费用相加 = 总费用 = 逐日相加
    t = payload["meta"]["totals"]
    check("汇总：三项费用之和 vs 总费用",
          abs(t["plan_cost_yuan"] + t["adjust_cost_yuan"] + t["emergency_cost_yuan"] - t["total_cost_yuan"]), 1e-6)
    check("汇总：逐日 total_cost 之和 vs 汇总",
          abs(sum(day["total_cost_yuan"] for day in days) - t["total_cost_yuan"]), 1e-6)
    check("汇总：逐日 total_cost 之和 vs 求解器 day_cost4",
          abs(sum(day["total_cost_yuan"] for day in days)
              - sum(M.day_cost4(first + i, x[first + i], q[first + i], {"z": z[first + i]})[0]
                    for i in range(len(days)))), 1e-6)

    print()
    print("=" * 88)
    if failures:
        print(f"验收未通过，{len(failures)} 项异常：")
        for item in failures:
            print("  -", item)
        return 1
    print(f"验收通过（变体 4-{which}）：物理约束、费用口径与工作簿逐格对账全部一致。")
    print("=" * 88)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
